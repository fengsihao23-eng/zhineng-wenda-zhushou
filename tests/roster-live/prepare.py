"""Prepare dictionaries from the supplied workbook and a separate destructive-test school.

Private fixture output is restricted to .local/yjyz-roster. All real staff stay
in the source school; the second school uses synthetic names and accounts.
"""
import asyncio
import hashlib
import json
import os
from pathlib import Path
from secrets import token_urlsafe
from uuid import uuid4
import httpx
from openpyxl import load_workbook
from sqlalchemy import select
from app.core.config import settings
from app.core.database import AsyncSessionLocal
from app.core.security import get_password_hash
from app.db.models import School, User, Role, UserRole
from app.services.roster_excel import TEACHER_COLUMNS, STUDENT_COLUMNS, GRADE_NAMES, teaching_pairs, workbook_bytes, text

ROOT = Path(__file__).resolve().parents[2]
PRIVATE = ROOT / '.local/yjyz-roster'
P = '/api/v1/platform/education'


async def call(client, method, path, body=None, token=None):
    response = await client.request(method, path, json=body, headers={'Authorization': 'Bearer ' + token, 'Idempotency-Key': str(uuid4())} if token else {})
    if not response.is_success:
        raise RuntimeError(f'{method} {path}: HTTP {response.status_code} / {response.headers.get("x-error-code", "unknown")}')
    return response.json()


async def main():
    if settings.DATABASE_URL != 'postgresql+asyncpg://qa_education@127.0.0.1:55449/yjyz_roster_20260922':
        raise RuntimeError('Dedicated roster database required')
    state = json.loads((PRIVATE / 'status.json').read_text())
    identity = json.loads((PRIVATE / 'identity.json').read_text())
    source = ROOT / '阳江一中高中教师资料模板0511.xlsx'
    book = load_workbook(source, read_only=True, data_only=False)
    rows = list(book.worksheets[0].iter_rows(values_only=True)); book.close()
    assert list(rows[0]) == TEACHER_COLUMNS
    original = [{k: text(v) for k,v in zip(TEACHER_COLUMNS, row)} for row in rows[1:] if any(v is not None for v in row)]
    pairs = {p for v in original for p in teaching_pairs(v['任课年级班级'])}
    classrooms = {(g,c) for _,g,c in pairs}
    for values in original:
        value = values['班主任年级班级']
        if value:
            grade,number=value.split('.')
            classrooms.add((int(grade),int(number)))
    subjects = sorted({s for s,_,_ in pairs})
    suffix = uuid4().hex[:10]
    password = token_urlsafe(32)
    async with AsyncSessionLocal() as db:
        school = School(id=uuid4(), code='QA-YJYZ-' + suffix, name='师生流程验收副本', source_system='QA')
        db.add(school); await db.flush()
        admin_role = await db.scalar(select(Role).where(Role.code=='SCHOOL_ADMIN'))
        admin=User(id=uuid4(),school_id=school.id,username='qa-roster-admin-'+suffix,account_type='general',display_name='流程验收管理员',password_hash=get_password_hash(password))
        db.add(admin); await db.flush()
        db.add(UserRole(user_id=admin.id,school_id=school.id,role_id=admin_role.id))
        await db.commit()
        sandbox={'school_id':str(school.id),'username':admin.username,'password':password}
    async with httpx.AsyncClient(base_url=state['baseUrl'],timeout=90,trust_env=False) as client:
        tokens=[]
        for actor in [identity,sandbox]:
            auth=await call(client,'POST','/api/v1/auth/login',{k:actor[k] for k in ['username','password','school_id']})
            token=auth['access_token']; tokens.append(token)
            present_classes=await call(client,'GET',P+'/school/classes?page_size=100',token=token)
            present_subjects=await call(client,'GET',P+'/school/subjects?page_size=100',token=token)
            existing_classes={r['external_class_id'] for r in present_classes['items']}
            existing_subjects={r['name'] for r in present_subjects['items']}
            for grade,number in sorted(classrooms):
                if f'{grade}.{number}' not in existing_classes:
                    await call(client,'POST',P+'/school/classes',{'external_class_id':f'{grade}.{number}','name':f'{GRADE_NAMES[grade]}{number}班'},token)
            for index,name in enumerate(subjects):
                if name not in existing_subjects:
                    await call(client,'POST',P+'/school/subjects',{'external_subject_id':f'YJYZ_SUBJECT_{index+1}','code':f'YJYZ_{index+1}','name':name},token)
        sandbox['token']=tokens[1]
    teachers=[]
    choices=[('teacher',0),('multi',157),('admin',154),('exam',155),('grade',158)]
    accounts={}
    for key,index in choices:
        values={**original[index], '教师姓名': '联调'+key, '教师账号':f'qa-{key}-{suffix}', '手机号码':'131****0000'}
        teachers.append(values); accounts[key]=values['教师账号']
    detached={c:'' for c in TEACHER_COLUMNS}; detached.update({'教师姓名':'联调无关联教师','教师账号':'qa-detached-'+suffix,'状态':'正常'})
    teachers.append(detached); accounts['detached']=detached['教师账号']
    students=[{'姓名':'联调学生甲' if i==0 else '联调学生乙','账号':f'qa-student-{i}-{suffix}','年级（1-12）':'高中一年级','班级号':'1班','状态':'正常'} for i in range(2)]
    files={
        'teachers':(TEACHER_COLUMNS,teachers),
        'students':(STUDENT_COLUMNS,students),
        'invalid':(TEACHER_COLUMNS,[{**detached,'教师姓名':'','教师账号':'','状态':'离职'}]),
        'suspects':(TEACHER_COLUMNS,[{**teachers[0],'教师账号':f'qa-suspect-{i}-{suffix}'} for i in range(2)]),
        'reimport':(TEACHER_COLUMNS,[{**detached,'教师姓名':'联调变更教师'}]),
        'teacher-change':(TEACHER_COLUMNS,[{**teachers[0],'教师姓名':'联调移交后重导教师'}]),
        'student-change':(STUDENT_COLUMNS,[{**students[1],'班级号':'2班'}]),
    }
    paths={}
    for key,(columns,data) in files.items():
        destination=PRIVATE/f'{key}-{suffix}.xlsx'
        destination.write_bytes(workbook_bytes(columns,[[v.get(c,'') for c in columns] for v in data]));destination.chmod(0o600)
        paths[key]=str(destination)
    fixture={'baseUrl':state['baseUrl'],'identity':identity,'sandbox':sandbox,'source':str(source),'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'original':original,'class_count':len(classrooms),'subject_count':len(subjects),'teaching_ref_count':len(pairs),'files':paths,'accounts':accounts,'students':students,'suffix':suffix}
    with os.fdopen(os.open(PRIVATE/'acceptance-fixture.json',os.O_WRONLY|os.O_CREAT|os.O_TRUNC,0o600),'w') as handle: json.dump(fixture,handle,ensure_ascii=False)
    print(json.dumps({'prepared':True,'teacher_rows':len(original),'classes':len(classrooms),'subjects':len(subjects),'sandbox':'isolated synthetic school'},ensure_ascii=False))


if __name__=='__main__': asyncio.run(main())
