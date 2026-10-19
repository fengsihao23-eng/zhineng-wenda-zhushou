"""Transactional v17 exam administration. Previewing never changes business data."""
import base64
import hashlib
import io
import json
import zipfile
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from uuid import UUID, uuid4
from sqlalchemy import select, delete, func
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
import xlrd

from app.core.errors import ApiError
from app.db.models import (School, User, Student, Subject, Exam, ExamSubject, StudentYearClass,
                           StudentExamPlacement, StudentSubjectScore, StudentExamScore, QuestionScore,
                           DiagnosisReport, TeacherProfile, TeachingAssignment, Role, UserRole)
from app.db.models.exam_management import (ExamAdministration, ExamParticipant, ExamQueryPolicy,
                                          ExamQueryGrant, ExamScoreFile, ExamOperation)
from app.services.education_common import data, digest, owned, claim, audit, now
from app.services import student_roster
from app.services.roster_excel import workbook_bytes
from app.services.exam_permission_defaults import DEFAULT_RULES, teacher_defaults

SCORE_TYPES = {"student": "个人单科成绩", "classSubject": "班级单科成绩", "gradeSubject": "年级单科成绩",
               "classAll": "班级综合成绩", "gradeAll": "年级综合成绩", "school": "全校综合成绩"}
DIAG_TYPES = {"student": "个人学情诊断", "classAll": "班级综合诊断", "classSubject": "班级单科诊断",
              "gradeAll": "年级综合诊断", "gradeSubject": "年级单科诊断", "school": "全校学情诊断"}
ESSAY_TYPES = {"chinese": "语文个人作文批改报告", "english": "英语个人作文批改报告"}
FIXED = ["序号", "考生号", "姓名", "行政班", "学校", "准考证号", "分数", "客观分", "主观分"]
LEGACY = ["考试ID", "学生ID", "学生姓名", "科目", "成绩"]
TEACHER_HEADERS = ["考试ID", "教师ID", "姓名", "成绩类型", "成绩科目", "诊断类型", "诊断科目", "作文报告", "学生范围"]
STUDENT_HEADERS = ["考试ID", "学生ID", "姓名", "年级", "班级（年级＋班级）", "操作", "操作科目", "变更前查询状态", "变更后查询范围"]


def fail(message, code="EXAM_INVALID", status=422):
    raise ApiError(status, code, message)


def unique(values):
    return list(dict.fromkeys(map(str, values)))


def class_key(student):
    return student["grade"] + " / " + student["class_number"]


async def options(db, actor, year=None):
    school = await student_roster.school(db, actor.school_id)
    current = student_roster.current_year(school)
    records = (await db.scalars(select(StudentYearClass).where(StudentYearClass.school_id == actor.school_id))).all()
    years = sorted({current, *(r.academic_year for r in records)}, reverse=True)
    year = year or current
    if year not in years:
        fail("请从系统已有学年中选择考试学年。")
    historic = {r.student_id: r for r in records if r.academic_year == year}
    classes = await student_roster.class_map(db, actor.school_id)
    students = []
    for student, user in await student_roster.people(db, actor.school_id):
        if student.status != "active":
            continue
        record = historic.get(student.id)
        if not record and year != current:
            continue
        view = student_roster.person_view(student, user, school, classes)
        if record:
            if record.status != "active":
                continue
            grade, number = record.grade, record.class_number
        else:
            grade, number = view["grade"], view["class_number"]
        if not grade or not number:
            continue
        students.append({"id": str(student.id), "platform_id": str(student.platform_id), "name": student.name,
                         "account": user.username if user else "", "grade": grade, "class_number": number,
                         "class_id": str(record.class_id if record else student.class_id) if (record.class_id if record else student.class_id) else None,
                         "external_class_id": student.external_class_id, "academic_year": year, "school": school.name})
    return {"current_year": current, "academic_year": year, "years": years, "school": school.name,
            "students": students, "score_types": SCORE_TYPES, "diagnosis_types": DIAG_TYPES, "essay_types": ESSAY_TYPES}


async def teacher_options(db, actor):
    teacher_ids = select(UserRole.user_id).join(Role, Role.id == UserRole.role_id).where(
        UserRole.school_id == actor.school_id, Role.code == "TEACHER")
    rows = (await db.execute(select(User, TeacherProfile).outerjoin(TeacherProfile, TeacherProfile.id == User.id).where(
        User.school_id == actor.school_id, User.status == "active", User.id.in_(teacher_ids)).order_by(User.display_name, User.id))).all()
    current_time = now()
    assignments = (await db.scalars(select(TeachingAssignment).where(TeachingAssignment.school_id == actor.school_id,
        TeachingAssignment.status == "active", TeachingAssignment.starts_at <= current_time,
        (TeachingAssignment.expires_at.is_(None) | (TeachingAssignment.expires_at > current_time))))).all()
    return [{"id": str(user.id), "platform_id": str(user.platform_id), "name": user.display_name or user.username,
             "account": user.username, "fields": profile.fields if profile else {}, "duties": profile.duties if profile else ["TEACHER"],
             "class_ids": unique(a.class_id for a in assignments if a.teacher_user_id == user.id),
             "subject_ids": unique(a.subject_id for a in assignments if a.teacher_user_id == user.id),
             "assignments": [{"class_id": str(assignment.class_id), "subject_id": str(assignment.subject_id)}
                             for assignment in assignments if assignment.teacher_user_id == user.id]} for user, profile in rows]


async def teacher_permission_options(db, actor, identifier):
    item = await administration(db, actor, identifier)
    ready = {str(file.subject_id) for file in await files_for(db, actor, identifier)}
    teachers = await teacher_options(db, actor)
    return {"items": [{**teacher, **teacher_defaults(teacher, item.configuration, ready)} for teacher in teachers],
            "default_rules": DEFAULT_RULES}


def teacher_default_rules_file():
    headers = ["教师角色", "默认成绩类型", "默认诊断类型", "默认科目", "默认学生范围"]
    rows = [[rule["label"], "、".join(SCORE_TYPES[value] for value in rule["types"]),
             "、".join(DIAG_TYPES[value] for value in rule["types"]), rule["subjects"], rule["scope"]]
            for rule in DEFAULT_RULES]
    book = load_workbook(io.BytesIO(workbook_bytes(headers, rows)))
    rules = book.active
    rules.title = "默认选项规则"
    for column, width in zip("ABCDE", [20, 44, 44, 28, 28]):
        rules.column_dimensions[column].width = width
    notes = book.create_sheet("使用说明")
    for values in [
        ["事项", "说明"],
        ["使用时机", "首次勾选未授权教师时生成默认选项；已保存、上传或手动编辑的配置优先保留。"],
        ["兼任角色", "多个角色的类型、科目和学生范围取并集。"],
        ["成绩与诊断", "成绩科目只默认勾选已导入科目；诊断科目限本次已配置科目，是否能查看具体报告取决于实际报告文件。"],
        ["作文报告", "作文批改报告默认不勾选，可按需手动选择。"],
        ["学生范围", "仅匹配本次考试的参试名单；无法匹配时请核对教师档案并手动补选。"],
        ["教师筛选", "班级筛选只匹配实际任课或班主任班级，默认授权范围不用于扩大筛选结果。年级、班级、学科同时筛选时须匹配同一档案职责。"],
        ["保存生效", "默认选项可手动调整，确认保存后才生效。下载本规则文件或授权预览均不会修改查询权限。"],
    ]:
        notes.append(values)
    notes.column_dimensions["A"].width = 20
    notes.column_dimensions["B"].width = 100
    notes.freeze_panes = "A2"
    for sheet in book:
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                if cell.row == 1:
                    cell.fill = PatternFill("solid", fgColor="5F5BDB")
                    cell.font = Font(bold=True, color="FFFFFF")
            sheet.row_dimensions[row[0].row].height = 30 if row[0].row == 1 else 84 if sheet == rules else 48
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


async def administration(db, actor, identifier, lock=False, revision=None):
    item = await owned(db, ExamAdministration, identifier, actor, lock=lock)
    if revision is not None and item.revision != revision:
        fail("考试配置或权限已更新，请刷新并重新预览后提交。", "EXAM_VERSION_CONFLICT", 409)
    return item


def roster_ids(config):
    return sorted({sid for subject in config["subjects"] for sid in subject["student_ids"]})


async def files_for(db, actor, identifier):
    return (await db.scalars(select(ExamScoreFile).where(ExamScoreFile.school_id == actor.school_id,
        ExamScoreFile.exam_id == identifier, ExamScoreFile.status == "imported"))).all()


async def detail(db, actor, identifier):
    item = await administration(db, actor, identifier)
    exam = await owned(db, Exam, identifier, actor)
    files = await files_for(db, actor, identifier)
    policies = (await db.scalars(select(ExamQueryPolicy).where(ExamQueryPolicy.school_id == actor.school_id,
        ExamQueryPolicy.exam_id == identifier))).all()
    teachers = {str(p.recipient_id): p.configuration for p in policies if p.recipient_kind == "teacher"}
    students = {str(p.recipient_id): p.configuration for p in policies if p.recipient_kind == "student" and p.configuration.get("subjects")}
    ids = roster_ids(item.configuration)
    required = {sid: [s["subject_id"] for s in item.configuration["subjects"] if sid in s["student_ids"]] for sid in ids}
    all_open = bool(ids) and all(set(required[sid]) <= set(students.get(sid, {}).get("subjects", [])) for sid in ids)
    state = "open" if teachers and all_open else "part" if teachers or students else "none"
    current_subjects = {s["subject_id"] for s in item.configuration["subjects"]}
    return {"id": str(exam.id), "exam_id": exam.external_exam_id or str(exam.id), "revision": item.revision,
            **item.configuration, "participant_count": len(ids), "teacher_permissions": teachers, "student_permissions": students,
            "open_state": state, "open_state_label": {"open": "已开放", "part": "部分开放", "none": "未开放"}[state],
            "score_files": {str(f.subject_id): {**data(f, "id", "filename", "created_at"), "row_count": len(f.rows)} for f in files if str(f.subject_id) in current_subjects}}


async def listing(db, actor, search="", year="", state="", grade="", subject="", page=1, page_size=20):
    query = select(ExamAdministration.id).join(Exam, Exam.id == ExamAdministration.id).where(ExamAdministration.school_id == actor.school_id)
    if search:
        query = query.where(Exam.name.contains(search, autoescape=True) | Exam.external_exam_id.contains(search, autoescape=True))
    if year:
        query = query.where(Exam.academic_year == year)
    # Configuration filters and open-state derive from the persisted snapshot.
    items = [await detail(db, actor, identifier) for identifier in (await db.scalars(query.order_by(Exam.created_at.desc(), Exam.id))).all()]
    grades = sorted({s["grade"] for item in items for s in item["snapshots"].values()})
    subjects = sorted({s["name"] for item in items for s in item["subjects"]})
    items = [item for item in items if (not state or item["open_state"] == state)
             and (not grade or any(s["grade"] == grade for sid, s in item["snapshots"].items() if sid in roster_ids(item)))
             and (not subject or any(s["name"] == subject for s in item["subjects"]))]
    return {"items": items[(page-1)*page_size:page*page_size], "total": len(items), "page": page, "page_size": page_size,
            "grades": grades, "subjects": subjects}


async def preview(db, actor, body, identifier=None):
    if identifier and body.expected_revision is None:
        fail("调整考试必须携带当前版本，请刷新后重新预览。")
    old = await administration(db, actor, identifier, revision=body.expected_revision) if identifier else None
    same_year = old is not None and body.academic_year == old.configuration["academic_year"]
    candidates = await options(db, actor, body.academic_year)
    available = {s["id"]: s for s in candidates["students"]}
    snapshots = dict(old.configuration["snapshots"]) if same_year else {}
    subjects = []
    old_by_id = {s["subject_id"]: s for s in old.configuration["subjects"]} if old else {}
    for subject in body.subjects:
        if subject.subject_id:
            if old and str(subject.subject_id) not in old_by_id:
                fail("科目标识不属于当前考试。")
            await owned(db, Subject, subject.subject_id, actor)
        ids = unique(subject.student_ids)
        retained = old_by_id.get(str(subject.subject_id), {}).get("student_ids", []) if same_year else []
        for sid in ids:
            if sid not in available and sid not in retained:
                fail("名单包含非本校、非考试学年或状态非正常的学生，请重新选择。")
            if sid not in snapshots:
                snapshots[sid] = available[sid]
        subjects.append({"key": subject.key, "subject_id": str(subject.subject_id) if subject.subject_id else None,
                         "name": subject.name, "full_score": float(subject.full_score), "student_ids": ids})
    duplicates = (await db.scalars(select(Exam).where(Exam.school_id == actor.school_id, Exam.academic_year == body.academic_year,
        Exam.name == body.name, Exam.id != identifier if identifier else True).order_by(Exam.id))).all()
    config = {"name": body.name, "academic_year": body.academic_year, "exam_date": body.exam_date.isoformat(),
              "subjects": subjects, "snapshots": snapshots}
    summary = defaultdict(lambda: defaultdict(int))
    for subject in subjects:
        for sid in subject["student_ids"]:
            summary[class_key(snapshots[sid])][subject["name"]] += 1
    changes = []
    if old:
        for s in subjects:
            prior = old_by_id.get(s["subject_id"])
            if prior != s:
                changes.append({"subject": s["name"], "before": prior, "after": s,
                                "added": sorted(set(s["student_ids"]) - set((prior or {}).get("student_ids", []))),
                                "removed": sorted(set((prior or {}).get("student_ids", [])) - set(s["student_ids"]))})
        for prior in old.configuration["subjects"]:
            if prior["subject_id"] not in {s["subject_id"] for s in subjects}:
                changes.append({"subject": prior["name"], "before": prior, "after": None, "removed": prior["student_ids"], "added": []})
    linked = {"score_records": int(await db.scalar(select(func.count()).select_from(StudentSubjectScore).where(StudentSubjectScore.exam_id == identifier)) or 0),
              "reports": int(await db.scalar(select(func.count()).select_from(DiagnosisReport).where(DiagnosisReport.exam_id == identifier)) or 0)} if old else {}
    if old:
        from app.db.models.report_management import ManagedReport
        linked["reports"] += int(await db.scalar(select(func.count()).select_from(ManagedReport).where(
            ManagedReport.school_id == actor.school_id, ManagedReport.exam_id == identifier, ManagedReport.status == "current")) or 0)
    checksum = digest({"config": config, "revision": old.revision if old else None, "duplicates": [str(e.id) for e in duplicates], "linked": linked})
    return {"configuration": config, "preview_hash": checksum, "participant_count": len(roster_ids(config)),
            "duplicates": [{"id": e.external_exam_id or str(e.id), "name": e.name} for e in duplicates],
            "summary": dict(summary), "changes": changes, "linked_data": linked,
            "permission_note": ("考试学年和本次名单归属按新学年更新。已有成绩、报告保留；原授权名单保持冻结，新增学生或科目须重新授权。"
                                if old and not same_year else "已有成绩、报告及历史归属保留。授权名单按授权时冻结，新增学生或科目须重新授权。")}


async def operation(db, actor, identifier, kind, details):
    user = await owned(db, User, actor.user_id, actor)
    item = ExamOperation(school_id=actor.school_id, exam_id=identifier, kind=kind, actor_id=actor.user_id,
                         operator_name=user.display_name or user.username, details=details)
    db.add(item)
    return item


async def save(db, actor, body, key=None, identifier=None):
    previous = await administration(db, actor, identifier, lock=True, revision=body.expected_revision) if identifier else None
    if not identifier:
        identifier, fresh = await claim(db, actor, "exam.v17.create", key, body.model_dump(mode="json"))
        if not fresh:
            return await detail(db, actor, identifier)
    checked = await preview(db, actor, body, previous.id if previous else None)
    if not body.preview_hash or checked["preview_hash"] != body.preview_hash:
        fail("考试或参试资料已变化，请重新预览并核对最新名单。", "EXAM_PREVIEW_CHANGED", 409)
    config = checked["configuration"]
    before = previous.configuration if previous else None
    if previous:
        exam = await owned(db, Exam, identifier, actor)
        exam.name, exam.academic_year, exam.start_date = body.name, body.academic_year, body.exam_date
    else:
        exam = Exam(id=identifier, school_id=actor.school_id, name=body.name, academic_year=body.academic_year,
                    start_date=body.exam_date, exam_type="school", source_system="native", status="active",
                    external_exam_id=f"EX-{body.exam_date:%Y%m%d}-{str(identifier)[:8].upper()}")
        db.add(exam)
        await db.flush()
    for subject in config["subjects"]:
        if not subject["subject_id"]:
            existing = await db.scalar(select(Subject).where(Subject.school_id == actor.school_id, Subject.name == subject["name"]).order_by(Subject.id).limit(1))
            if existing is None:
                existing = Subject(id=uuid4(), school_id=actor.school_id, name=subject["name"], code="EX-"+uuid4().hex[:16], source_system="native")
                db.add(existing)
                await db.flush()
            subject["subject_id"] = str(existing.id)
        configured = await db.scalar(select(ExamSubject).where(ExamSubject.exam_id == identifier, ExamSubject.subject_id == UUID(subject["subject_id"])))
        if configured:
            configured.full_score = Decimal(str(subject["full_score"]))
        else:
            db.add(ExamSubject(exam_id=identifier, subject_id=UUID(subject["subject_id"]), full_score=Decimal(str(subject["full_score"]))))
    if len({s["subject_id"] for s in config["subjects"]}) != len(config["subjects"]):
        fail("所选科目实际指向同一学科，请核对科目名称。")
    await db.execute(delete(ExamParticipant).where(ExamParticipant.school_id == actor.school_id, ExamParticipant.exam_id == identifier))
    for subject in config["subjects"]:
        for sid in subject["student_ids"]:
            db.add(ExamParticipant(school_id=actor.school_id, exam_id=identifier, subject_id=UUID(subject["subject_id"]), student_id=UUID(sid)))
    existing_placements = {placement.student_id: placement for placement in (await db.scalars(select(StudentExamPlacement).where(
        StudentExamPlacement.school_id == actor.school_id, StudentExamPlacement.exam_id == identifier))).all()}
    for sid in roster_ids(config):
        placement = existing_placements.get(UUID(sid))
        if placement is None:
            placement = StudentExamPlacement(school_id=actor.school_id, exam_id=identifier, student_id=UUID(sid))
            db.add(placement)
        if UUID(sid) not in existing_placements or body.academic_year != before["academic_year"]:
            snap = config["snapshots"][sid]
            placement.academic_year, placement.grade, placement.class_number = body.academic_year, snap["grade"], snap["class_number"]
            placement.class_id = UUID(snap["class_id"]) if snap.get("class_id") else None
            placement.external_class_id = snap.get("external_class_id")
    if previous:
        previous.configuration = config
        previous.revision += 1
    else:
        db.add(ExamAdministration(id=identifier, school_id=actor.school_id, configuration=config, revision=1))
    await operation(db, actor, identifier, "adjust" if previous else "create", {"before": before, "after": config,
        "changes": checked["changes"], "linked_data": checked["linked_data"], "permissions_expanded": False})
    audit(db, actor, "exam.adjust" if previous else "exam.create", exam)
    await db.commit()
    return await detail(db, actor, identifier)


def decode_upload(body):
    if not body.filename.lower().endswith((".xlsx", ".xls")):
        fail("请选择 Excel 文件（.xlsx 或 .xls）。")
    try:
        content = base64.b64decode(body.content_base64, validate=True)
        if not content or len(content) > 10_000_000:
            fail("Excel 文件为空或超过 10 MB。")
        return content
    except (ValueError, TypeError):
        fail("文件编码无效，请重新选择 Excel。")


def matrix(body):
    content = decode_upload(body)
    try:
        if body.filename.lower().endswith(".xls"):
            book = xlrd.open_workbook(file_contents=content, on_demand=True)
            sheet = book.sheet_by_index(0)
            if sheet.nrows > 30010 or sheet.ncols > 509:
                fail("Excel 超过 30000 行或 500 个小题列。")
            result = [(i+1, sheet.row_values(i)) for i in range(sheet.nrows)]
            book.release_resources()
        else:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 60_000_000:
                    fail("Excel 解压后过大，请拆分文件。")
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            sheet = book.worksheets[0]
            if (sheet.max_row or 0) > 30010 or (sheet.max_column or 0) > 509:
                fail("Excel 超过 30000 行或 500 个小题列。")
            result = []
            for i, row in enumerate(sheet.iter_rows(), 1):
                if any(cell.data_type == "f" for cell in row):
                    fail(f"第 {i} 行包含公式，请粘贴为数值后上传。")
                result.append((i, [c.value if c.value is not None else "" for c in row]))
            book.close()
        return content, [(line, [str(int(v)) if isinstance(v, float) and v.is_integer() else str(v).strip() for v in row])
                         for line, row in result if any(str(v).strip() for v in row)]
    except ApiError:
        raise
    except Exception:
        fail("Excel 无法读取，请使用未加密且未损坏的工作簿。")


def number(value, label, allow_blank=False):
    if value == "" and allow_blank:
        return None
    try:
        result = Decimal(str(value))
        if not result.is_finite() or result < 0 or result > Decimal("1000000") or result.as_tuple().exponent < -2:
            raise ValueError()
        return result
    except (ValueError, InvalidOperation):
        raise ValueError(f"{label}须为最多两位小数的非负数") from None


def subject_config(config, subject_id):
    result = next((s for s in config["subjects"] if s["subject_id"] == str(subject_id)), None)
    if result is None:
        fail("科目不属于本次考试的当前配置。")
    return result


def score_validation(exam, config, subject, rows):
    errors, parsed, seen = [], [], set()
    lookup = {str(s["platform_id"]): sid for sid, s in config["snapshots"].items()}
    expected = set(subject["student_ids"])
    if not rows:
        return [], ["工作簿为空。"]
    legacy = rows[0][1] == LEGACY
    header_index = 0 if legacy else next((i for i, (_, row) in enumerate(rows[:6]) if row[:9] == FIXED), -1)
    if header_index < 0:
        return [], ["表头前 9 列须依次为：" + " / ".join(FIXED)]
    header = rows[header_index][1]
    questions = [] if legacy else header[9:]
    if any(not q or len(q) > 20 for q in questions) or len(set(questions)) != len(questions):
        return [], ["小题题号不得为空、重复或超过20个字符。"]
    for line, raw in rows[header_index+1:]:
        row = raw + [""] * max(0, len(header)-len(raw))
        try:
            if len(row) > len(header) and any(row[len(header):]):
                raise ValueError("存在表头以外的多余数据列")
            if legacy:
                eid, sid_text, name, subject_name, total = row[:5]
                if eid not in {str(exam.id), exam.external_exam_id}:
                    raise ValueError("考试ID不符")
                if subject_name != subject["name"]:
                    raise ValueError("科目与当前导入科目不符")
                account, classroom, school_name = "", "", ""
                qvalues, objective, subjective = {}, None, None
            else:
                _, account, name, classroom, school_name, sid_text, total, obj, sub = row[:9]
                objective, subjective = number(obj, "客观分", True), number(sub, "主观分", True)
                qvalues = {q: float(number(row[9+i], f"小题「{q}」得分")) for i, q in enumerate(questions) if row[9+i] != ""}
                if not total and qvalues:
                    total = str(sum(Decimal(str(v)) for v in qvalues.values()))
            sid = lookup.get(sid_text, sid_text)
            if sid not in expected:
                raise ValueError("准考证号（学生ID）不在该科参试名单")
            snap = config["snapshots"][sid]
            if name != snap["name"]:
                raise ValueError("姓名与考试名单不符")
            if account and account != snap["account"]:
                raise ValueError("考生号（账号）与考试名单不符")
            if school_name and school_name != snap["school"]:
                raise ValueError("学校与本校不符")
            if classroom and classroom not in {snap["class_number"], class_key(snap), snap["grade"]+snap["class_number"]}:
                raise ValueError("行政班与考试学年名单不符")
            value = number(total, "分数")
            if value > Decimal(str(subject["full_score"])):
                raise ValueError("分数超过科目满分")
            if any(Decimal(str(v)) > Decimal(str(subject["full_score"])) for v in qvalues.values()):
                raise ValueError("小题得分超过科目满分")
            if sid in seen:
                raise ValueError("学生重复（学生ID重复）")
            seen.add(sid)
            parsed.append({"student_id": sid, "score": float(value), "objective": float(objective) if objective is not None else None,
                           "subjective": float(subjective) if subjective is not None else None, "questions": qvalues})
        except ValueError as exc:
            errors.append(f"第 {line} 行：{exc}")
    missing = expected - seen
    if missing:
        errors.append(f"必须导入完整名单，缺少 {len(missing)} 名参试学生。")
    if len(parsed) != len(expected):
        errors.append(f"成绩行数须与本次该科完整参试名单一致（{len(expected)} 人）。")
    return parsed, errors


async def score_template(db, actor, identifier, subject_id):
    item = await administration(db, actor, identifier)
    subject = subject_config(item.configuration, subject_id)
    snaps = item.configuration["snapshots"]
    return workbook_bytes(FIXED, [[i, snaps[sid]["account"], snaps[sid]["name"], snaps[sid]["class_number"],
        snaps[sid]["school"], snaps[sid]["platform_id"], "", "", ""] for i, sid in enumerate(subject["student_ids"], 1)])


async def score_preflight(db, actor, identifier, body):
    item = await administration(db, actor, identifier, lock=True, revision=body.expected_revision)
    subject = subject_config(item.configuration, body.subject_id)
    if any(f.subject_id == body.subject_id for f in await files_for(db, actor, identifier)):
        fail("该科已有成绩文件，请先在「查看已导入成绩」中删除后再导入。", "SCORE_FILE_EXISTS", 409)
    content, rows = matrix(body)
    exam = await owned(db, Exam, identifier, actor)
    parsed, errors = score_validation(exam, item.configuration, subject, rows)
    batch = ExamScoreFile(id=uuid4(), school_id=actor.school_id, exam_id=identifier, subject_id=body.subject_id,
        filename=body.filename, original_content=content, content_hash=hashlib.sha256(content).hexdigest(),
        rows=parsed if not errors else [], errors=errors, revision=item.revision, created_by=actor.user_id,
        status="invalid" if errors else "validated")
    db.add(batch)
    await db.commit()
    return {**data(batch, "id", "status", "filename", "revision", "errors"), "row_count": len(batch.rows)}


async def recalculate_totals(db, actor, identifier, config):
    rows = (await db.scalars(select(StudentSubjectScore).where(StudentSubjectScore.school_id == actor.school_id,
        StudentSubjectScore.exam_id == identifier))).all()
    by_student = defaultdict(list)
    for score in rows:
        by_student[score.student_id].append(score)
    previous = {s.student_id: s for s in (await db.scalars(select(StudentExamScore).where(StudentExamScore.school_id == actor.school_id,
        StudentExamScore.exam_id == identifier))).all()}
    # Totals are only materialized for students whose complete participation has scores.
    for sid, scores in by_student.items():
        required = {s["subject_id"] for s in config["subjects"] if str(sid) in s["student_ids"]}
        if not required or not required <= {str(s.subject_id) for s in scores}:
            if sid in previous:
                await db.delete(previous[sid])
            continue
        total = previous.get(sid)
        if total is None:
            total = StudentExamScore(school_id=actor.school_id, student_id=sid, exam_id=identifier)
            db.add(total)
        total.total_score = sum(s.score for s in scores)
        total.full_score = sum(s.full_score for s in scores)
    for sid, total in previous.items():
        if sid not in by_student:
            await db.delete(total)
    await db.flush()
    from app.services.student_data import recalculate, location
    placements = (await db.scalars(select(StudentExamPlacement).where(StudentExamPlacement.school_id == actor.school_id,
        StudentExamPlacement.exam_id == identifier))).all()
    classes, grades = set(), set()
    for p in placements:
        grade, class_number = location(p)
        for sub in [None] + list({s.subject_id for s in rows}):
            classes.add((identifier, sub, grade, class_number))
            grades.add((identifier, sub, grade))
    await recalculate(db, actor, {identifier}, classes, grades)


async def score_confirm(db, actor, identifier, file_id):
    batch = await owned(db, ExamScoreFile, file_id, actor, lock=True)
    if batch.exam_id != identifier:
        fail("成绩批次不属于本次考试。")
    item = await administration(db, actor, identifier, lock=True)
    if batch.status == "imported":
        return await detail(db, actor, identifier)
    if batch.status != "validated":
        fail("仅校验通过的文件可确认导入。")
    if batch.revision != item.revision:
        fail("预校验后考试已变化，请重新校验文件。", "EXAM_VERSION_CONFLICT", 409)
    if any(f.subject_id == batch.subject_id for f in await files_for(db, actor, identifier)):
        fail("该科已有成绩文件，请先删除后重新导入。", "SCORE_FILE_EXISTS", 409)
    sub = subject_config(item.configuration, batch.subject_id)
    for row in batch.rows:
        db.add(StudentSubjectScore(school_id=actor.school_id, exam_id=identifier, student_id=UUID(row["student_id"]),
            subject_id=batch.subject_id, score=Decimal(str(row["score"])), full_score=Decimal(str(sub["full_score"]))))
        for question, value in row["questions"].items():
            # Imported files provide earned scores only, never invent question maxima or losses.
            db.add(QuestionScore(school_id=actor.school_id, exam_id=identifier, student_id=UUID(row["student_id"]),
                subject_id=batch.subject_id, question_no=question, score=Decimal(str(value)), full_score=None, lost_score=None))
    batch.status = "imported"
    item.revision += 1
    await db.flush()
    await recalculate_totals(db, actor, identifier, item.configuration)
    await operation(db, actor, identifier, "score_import", {"file_id": str(batch.id), "filename": batch.filename,
        "subject_id": str(batch.subject_id), "subject_name": sub["name"], "students": batch.rows, "permissions_opened": False})
    await db.commit()
    return await detail(db, actor, identifier)


async def score_delete(db, actor, identifier, subject_id, body):
    item = await administration(db, actor, identifier, lock=True, revision=body.expected_revision)
    files = [f for f in await files_for(db, actor, identifier) if f.subject_id == subject_id]
    if not files:
        fail("该科没有已导入的成绩文件。", status=404)
    for model in (QuestionScore, StudentSubjectScore):
        await db.execute(delete(model).where(model.school_id == actor.school_id, model.exam_id == identifier, model.subject_id == subject_id))
    # Student grants cover scores and reports for that subject; teacher report grants are independent.
    await db.execute(delete(ExamQueryGrant).where(ExamQueryGrant.school_id == actor.school_id,
        ExamQueryGrant.exam_id == identifier, ExamQueryGrant.subject_id == subject_id,
        (ExamQueryGrant.recipient_kind == "student") | ExamQueryGrant.content_type.like("score:%")))
    policies = (await db.scalars(select(ExamQueryPolicy).where(ExamQueryPolicy.school_id == actor.school_id,
        ExamQueryPolicy.exam_id == identifier))).all()
    for policy in policies:
        config = dict(policy.configuration)
        field = "subjects" if policy.recipient_kind == "student" else "score_subjects"
        config[field] = [s for s in config.get(field, []) if s != str(subject_id)]
        if field == "score_subjects" and not config[field]:
            config["score_types"] = []
        if policy.recipient_kind == "teacher" and not (config.get("score_types") or config.get("diagnosis_types") or config.get("essay_types")):
            await db.delete(policy)
        else:
            policy.configuration = config
    for file in files:
        file.status = "deleted"
    item.revision += 1
    await db.flush()
    await recalculate_totals(db, actor, identifier, item.configuration)
    await operation(db, actor, identifier, "score_delete", {"subject_id": str(subject_id), "files": [f.filename for f in files],
        "subject_name": next((s["name"] for s in item.configuration["subjects"] if s["subject_id"] == str(subject_id)), "历史科目"),
        "removed_score_grants": True, "reports_preserved": True})
    await db.commit()
    return await detail(db, actor, identifier)


async def score_file(db, actor, identifier, file_id):
    file = await owned(db, ExamScoreFile, file_id, actor)
    if file.exam_id != identifier or file.status != "imported":
        fail("成绩文件不存在或已删除。", status=404)
    return file


async def logs(db, actor, identifier=None, category="", page=1, page_size=20):
    query = select(ExamOperation).where(ExamOperation.school_id == actor.school_id)
    if identifier:
        await administration(db, actor, identifier)
        query = query.where(ExamOperation.exam_id == identifier)
    if category == "permission":
        query = query.where(ExamOperation.kind.in_(["teacher_permission", "student_permission", "teacher_revoke"]))
    elif category == "exam":
        query = query.where(~ExamOperation.kind.in_(["teacher_permission", "student_permission", "teacher_revoke"]))
    total = int(await db.scalar(select(func.count()).select_from(query.subquery())) or 0)
    rows = (await db.scalars(query.order_by(ExamOperation.created_at.desc(), ExamOperation.id.desc()).offset((page-1)*page_size).limit(page_size))).all()
    exams = {e.id: e for e in (await db.scalars(select(Exam).where(Exam.school_id == actor.school_id, Exam.id.in_([r.exam_id for r in rows])))).all()}
    return {"items": [{**data(r, "id", "exam_id", "kind", "operator_name", "created_at", "details"),
        "exam_name": exams[r.exam_id].name, "exam_code": exams[r.exam_id].external_exam_id} for r in rows], "total": total}


def selected_range(config, permission):
    ids = roster_ids(config)
    if permission.range_type == "grade":
        ids = [sid for sid in ids if config["snapshots"][sid]["grade"] in permission.grades]
    elif permission.range_type == "class":
        ids = [sid for sid in ids if class_key(config["snapshots"][sid]) in permission.classes]
    if not ids:
        fail("请选择本次参试名单中的有效学生范围。")
    return ids


async def permission_preview(db, actor, identifier, body, kind):
    item = await administration(db, actor, identifier, revision=body.expected_revision)
    config = item.configuration
    ready = {str(f.subject_id) for f in await files_for(db, actor, identifier)}
    subject_names = {s["subject_id"]: s["name"] for s in config["subjects"]}
    grants, policies, export_rows = [], {}, []
    current = (await db.scalars(select(ExamQueryPolicy).where(ExamQueryPolicy.school_id == actor.school_id,
        ExamQueryPolicy.exam_id == identifier, ExamQueryPolicy.recipient_kind == kind))).all()
    previous = {str(p.recipient_id): p.configuration for p in current}
    exam = await owned(db, Exam, identifier, actor)
    if kind == "teacher":
        teachers = {t["id"]: t for t in await teacher_options(db, actor)}
        if len({p.teacher_id for p in body.teachers}) != len(body.teachers):
            fail("同一教师不能重复指定。")
        for permission in body.teachers:
            tid = str(permission.teacher_id)
            if tid not in teachers:
                fail("教师不属于本校有效教师账号。")
            selected = selected_range(config, permission)
            if not (permission.score_types or permission.diagnosis_types or permission.essay_types):
                fail("请选择成绩数据、诊断报告或作文批改报告。")
            if bool(permission.score_types) != bool(permission.score_subjects) or bool(permission.diagnosis_types) != bool(permission.diagnosis_subjects):
                fail("成绩和诊断报告均须同时选择类型与科目。")
            if set(unique(permission.score_subjects)) - ready:
                fail("未导入成绩的科目不能开放成绩查询。")
            if set(unique(permission.score_subjects)) - set(subject_names):
                fail("成绩科目不属于本次考试的当前配置。")
            if set(unique(permission.diagnosis_subjects)) - set(subject_names):
                fail("诊断报告科目不属于本次考试。")
            essay_subjects = {sid: "chinese" if name == "语文" else "english" for sid, name in subject_names.items()
                              if name in {"语文", "英语"} and ("chinese" if name == "语文" else "english") in permission.essay_types}
            if set(permission.essay_types) - set(essay_subjects.values()):
                fail("本次考试未配置所选语种的作文科目。")
            for subject in config["subjects"]:
                sid = subject["subject_id"]
                content = (["score:"+t for t in set(permission.score_types)] if sid in unique(permission.score_subjects) else [])
                content += (["diagnosis:"+t for t in set(permission.diagnosis_types)] if sid in unique(permission.diagnosis_subjects) else [])
                content += ["essay:"+essay_subjects[sid]] if sid in essay_subjects else []
                for student in set(subject["student_ids"]) & set(selected):
                    grants.extend({"recipient_id": tid, "student_id": student, "subject_id": sid, "content_type": c} for c in content)
            pconfig = permission.model_dump(mode="json")
            pconfig["frozen_student_ids"] = selected
            policies[tid] = pconfig
            scope = "全部参试学生" if permission.range_type == "all" else "年级："+"、".join(permission.grades) if permission.range_type == "grade" else "班级："+"、".join(permission.classes)
            export_rows.append([exam.external_exam_id, teachers[tid]["platform_id"], teachers[tid]["name"],
                "、".join(SCORE_TYPES[t] for t in permission.score_types), "、".join(subject_names[s] for s in unique(permission.score_subjects)),
                "、".join(DIAG_TYPES[t] for t in permission.diagnosis_types), "、".join(subject_names[s] for s in unique(permission.diagnosis_subjects)),
                "、".join(ESSAY_TYPES[t] for t in permission.essay_types), scope])
    else:
        ids = set(roster_ids(config))
        if set(unique(body.student_ids)) - ids:
            fail("学生不在本次考试参试名单中。")
        subjects = unique(body.subject_ids)
        if set(subjects) - set(subject_names):
            fail("所选科目已失效。")
        if body.mode == "open" and set(subjects) - ready:
            fail("未导入成绩的科目不能开放。")
        for student in unique(body.student_ids):
            before = set(previous.get(student, {}).get("subjects", []))
            applicable = {s["subject_id"] for s in config["subjects"] if student in s["student_ids"]} & set(subjects)
            if body.mode == "open":
                scored = set((await db.scalars(select(StudentSubjectScore.subject_id).where(StudentSubjectScore.school_id == actor.school_id,
                    StudentSubjectScore.exam_id == identifier, StudentSubjectScore.student_id == UUID(student)))).all())
                if applicable - {str(s) for s in scored}:
                    fail("所选学生的科目成绩尚未完整导入。")
            after = before | applicable if body.mode == "open" else before - set(subjects)
            policies[student] = {"subjects": sorted(after)}
            for sub in after:
                grants.append({"recipient_id": student, "student_id": student, "subject_id": sub, "content_type": "student"})
            snap = config["snapshots"][student]
            names = lambda values: "、".join(subject_names.get(s, "历史科目") for s in sorted(values)) or "未授权"
            export_rows.append([exam.external_exam_id, snap["platform_id"], snap["name"], snap["grade"], class_key(snap),
                "开放" if body.mode == "open" else "关闭", names(subjects), names(before), names(after)])
    grants.sort(key=lambda g: (g["recipient_id"], g["student_id"], g["subject_id"], g["content_type"]))
    checksum = digest({"exam": str(identifier), "revision": item.revision, "kind": kind, "policies": policies, "grants": grants})
    return {"preview_hash": checksum, "policies": policies, "grants": grants, "rows": export_rows,
            "headers": TEACHER_HEADERS if kind == "teacher" else STUDENT_HEADERS,
            "before": {key: previous.get(key) for key in policies}}


async def permission_save(db, actor, identifier, body, kind):
    item = await administration(db, actor, identifier, lock=True, revision=body.expected_revision)
    checked = await permission_preview(db, actor, identifier, body, kind)
    if body.preview_hash is not None and body.preview_hash != checked["preview_hash"]:
        fail("授权配置已变化，请核对当前选项后重新保存。", "EXAM_PREVIEW_CHANGED", 409)
    ids = [UUID(s) for s in checked["policies"]]
    await db.execute(delete(ExamQueryGrant).where(ExamQueryGrant.school_id == actor.school_id, ExamQueryGrant.exam_id == identifier,
        ExamQueryGrant.recipient_kind == kind, ExamQueryGrant.recipient_id.in_(ids)))
    await db.execute(delete(ExamQueryPolicy).where(ExamQueryPolicy.school_id == actor.school_id, ExamQueryPolicy.exam_id == identifier,
        ExamQueryPolicy.recipient_kind == kind, ExamQueryPolicy.recipient_id.in_(ids)))
    for recipient, configuration in checked["policies"].items():
        db.add(ExamQueryPolicy(school_id=actor.school_id, exam_id=identifier, recipient_kind=kind, recipient_id=UUID(recipient), configuration=configuration))
    for grant in checked["grants"]:
        db.add(ExamQueryGrant(school_id=actor.school_id, exam_id=identifier, recipient_kind=kind,
            **{key: UUID(value) if key.endswith("_id") else value for key, value in grant.items()}))
    item.revision += 1
    await operation(db, actor, identifier, kind+"_permission", {"before": checked["before"], "after": checked["policies"],
        "preview_rows": checked["rows"], "frozen_grants": checked["grants"]})
    await db.commit()
    return await detail(db, actor, identifier)


async def revoke_teacher(db, actor, identifier, teacher_id, body):
    item = await administration(db, actor, identifier, lock=True, revision=body.expected_revision)
    prior = await db.scalar(select(ExamQueryPolicy).where(ExamQueryPolicy.school_id == actor.school_id, ExamQueryPolicy.exam_id == identifier,
        ExamQueryPolicy.recipient_kind == "teacher", ExamQueryPolicy.recipient_id == teacher_id))
    if prior is None:
        fail("该教师尚未授权。")
    before = prior.configuration
    teacher = await owned(db, User, teacher_id, actor)
    await db.execute(delete(ExamQueryGrant).where(ExamQueryGrant.school_id == actor.school_id, ExamQueryGrant.exam_id == identifier,
        ExamQueryGrant.recipient_kind == "teacher", ExamQueryGrant.recipient_id == teacher_id))
    await db.delete(prior)
    item.revision += 1
    await operation(db, actor, identifier, "teacher_revoke", {"teacher_id": str(teacher_id), "teacher_platform_id": str(teacher.platform_id),
        "teacher_name": teacher.display_name or teacher.username, "before": before, "after": None})
    await db.commit()
    return await detail(db, actor, identifier)


async def permission_upload(db, actor, identifier, body, kind):
    from app.schemas.exam_management import TeacherPermission
    item = await administration(db, actor, identifier)
    _, rows = matrix(body)
    expected = TEACHER_HEADERS if kind == "teacher" else STUDENT_HEADERS
    header_index = next((i for i, (_, row) in enumerate(rows[:6]) if row == expected), -1)
    if header_index < 0:
        fail("请使用本页下载的新版授权预览表，列名与顺序须一致。")
    people = {t["platform_id"]: t["id"] for t in await teacher_options(db, actor)} if kind == "teacher" else {
        s["platform_id"]: sid for sid, s in item.configuration["snapshots"].items()}
    names = {s["name"]: s["subject_id"] for s in item.configuration["subjects"]}
    exam = await owned(db, Exam, identifier, actor)
    result, seen = [], set()
    def parse(value, mapping):
        reverse = {label: key for key, label in mapping.items()}
        try:
            return [reverse[v] for v in value.split("、") if v]
        except KeyError:
            fail("授权表包含无法识别的查询类型。")
    for line, raw in rows[header_index+1:]:
        if len(raw) != len(expected):
            fail(f"第 {line} 行列数不一致。")
        row = dict(zip(expected, raw))
        if row["考试ID"] not in {exam.external_exam_id, str(identifier)}:
            fail(f"第 {line} 行考试ID不符。")
        person = people.get(row["教师ID" if kind == "teacher" else "学生ID"])
        if not person or person in seen:
            fail(f"第 {line} 行人员ID无效或重复。")
        seen.add(person)
        if kind == "teacher":
            try:
                scope_text = row["学生范围"]
                if scope_text == "全部参试学生":
                    scope = {"type": "all"}
                elif scope_text.startswith("年级："):
                    scope = {"type": "grade", "grades": scope_text[3:].split("、")}
                elif scope_text.startswith("班级："):
                    scope = {"type": "class", "classes": scope_text[3:].split("、")}
                else:
                    scope = json.loads(scope_text)
                p = TeacherPermission(teacher_id=person, score_types=parse(row["成绩类型"], SCORE_TYPES),
                    score_subjects=[names[v] for v in row["成绩科目"].split("、") if v], diagnosis_types=parse(row["诊断类型"], DIAG_TYPES),
                    diagnosis_subjects=[names[v] for v in row["诊断科目"].split("、") if v], essay_types=parse(row["作文报告"], ESSAY_TYPES),
                    range_type=scope["type"], grades=scope.get("grades", []), classes=scope.get("classes", []))
            except (KeyError, ValueError, TypeError):
                fail(f"第 {line} 行权限或学生范围格式不正确。")
            result.append(p.model_dump(mode="json"))
        else:
            if row["操作"] not in {"开放", "关闭"}:
                fail(f"第 {line} 行操作须为开放或关闭。")
            try:
                subjects = [names[v] for v in row["操作科目"].split("、") if v]
            except KeyError:
                fail(f"第 {line} 行包含未知科目。")
            result.append({"student_id": person, "mode": "open" if row["操作"] == "开放" else "close", "subject_ids": subjects})
    if not result:
        fail("授权表没有人员数据。")
    return {"items": result, "revision": item.revision, "saved": False}
