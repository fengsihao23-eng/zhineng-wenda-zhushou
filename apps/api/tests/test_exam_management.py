"""Real HTTP/SQL acceptance for exam v17 creation, imports and data access."""
import base64
import io
import hashlib
from datetime import datetime, timezone
from uuid import UUID, uuid4
from types import SimpleNamespace
import pytest
from openpyxl import load_workbook, Workbook
from sqlalchemy import select, func
from app.db.models import (School, User, Student, StudentYearClass, StudentExamPlacement, StudentSubjectScore,
                           StudentExamScore, QuestionScore, Exam, DiagnosisReport, TeacherProfile)
from app.db.models.exam_management import ExamAdministration, ExamQueryGrant, ExamOperation
from app.db.models.education import PrivateAsset, ReportAttachment
from app.tools.exam_tools import GetExamSummaryTool
from app.tools.score_tools import GetSubjectScoresTool
from app.tools.base import ToolContext
from app.services.roster_excel import workbook_bytes
from test_education_workbench import actors  # noqa: F401

ROOT = "/api/v1/platform/education/exam-management"


@pytest.fixture
async def exam_people(test_db, actors):
    school = await test_db.get(School, actors["school"].id)
    school.current_academic_year = "2026-2027"
    first = await test_db.get(Student, actors["student"].id)
    first.name = "同名学生"
    rows = [first]
    for index in range(2):
        user = User(school_id=school.id, username="exam-student-"+str(index), password_hash="unused", display_name="合成学生")
        test_db.add(user)
        await test_db.flush()
        student = Student(school_id=school.id, user_id=user.id, name="同名学生" if index == 0 else "新增学生")
        test_db.add(student)
        await test_db.flush()
        rows.append(student)
    for student in rows:
        test_db.add(StudentYearClass(school_id=school.id, student_id=student.id, academic_year="2026-2027", grade="高中一年级", class_number="1班"))
    await test_db.commit()
    return [SimpleNamespace(id=s.id) for s in rows]


def switch(actors, db, role):
    db.info.pop("exam_student_scope", None)
    db.info.pop("exam_teacher_scope", None)
    actors["switch"](role)


def draft(people, count=2, subjects=2):
    return {"name": "合成考务验收", "academic_year": "2026-2027", "exam_date": "2026-10-16",
            "subjects": [{"key": str(i), "name": name, "full_score": 150, "student_ids": [str(s.id) for s in people[:count]]}
                         for i, name in enumerate(["数学", "语文"][:subjects])]}


async def post(client, path, body=None, status=200, key=None):
    response = await client.post(ROOT+path, json=body, headers={"Idempotency-Key": str(key or uuid4())})
    assert response.status_code == status, response.text
    return response.json()


async def create(client, people, **kwargs):
    body = draft(people, **kwargs)
    checked = await post(client, "/preview", body)
    return await post(client, "", {**body, "preview_hash": checked["preview_hash"]})


async def score_file(client, exam, subject_index=0, values=None, questions=False):
    subject = exam["subjects"][subject_index]
    response = await client.get(ROOT+f"/{exam['id']}/score-template/{subject['subject_id']}")
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    sheet = book.active
    for row in range(2, sheet.max_row+1):
        sheet.cell(row, 7, (values or [90, 95, 80])[row-2])
        if questions:
            sheet.cell(1, 10, "1"); sheet.cell(1, 11, "作文")
            sheet.cell(row, 10, 40); sheet.cell(row, 11, 50)
    output = io.BytesIO(); book.save(output); book.close()
    payload = {"subject_id": subject["subject_id"], "expected_revision": exam["revision"], "filename": "合成成绩.xlsx",
               "content_base64": base64.b64encode(output.getvalue()).decode()}
    return payload, output.getvalue()


async def import_score(client, exam, subject_index=0, questions=False):
    body, original = await score_file(client, exam, subject_index, questions=questions)
    checked = await post(client, f"/{exam['id']}/score-files", body)
    assert checked["status"] == "validated", checked
    result = await post(client, f"/{exam['id']}/score-files/{checked['id']}/confirm")
    return result, checked, original


async def permission(client, exam, kind, body):
    payload = {"expected_revision": exam["revision"], **body}
    checked = await client.post(ROOT+f"/{exam['id']}/permissions/{kind}/preview", json=payload)
    assert checked.status_code == 200, checked.text
    checksum = checked.headers["X-Exam-Preview-Hash"]
    return await post(client, f"/{exam['id']}/permissions/{kind}", {**payload, "preview_hash": checksum})


@pytest.mark.asyncio
async def test_create_preview_cancel_dedup_same_name_and_idempotency(client, test_db, actors, exam_people):
    body = draft(exam_people)
    body["subjects"][0]["student_ids"].append(str(exam_people[0].id))
    preview = await post(client, "/preview", body)
    assert preview["participant_count"] == 2
    assert await test_db.scalar(select(func.count()).select_from(ExamAdministration)) == 0
    key = uuid4()
    saved = await post(client, "", {**body, "preview_hash": preview["preview_hash"]}, key=key)
    repeated = await post(client, "", {**body, "preview_hash": preview["preview_hash"]}, key=key)
    assert saved["id"] == repeated["id"] and saved["participant_count"] == 2
    assert saved["open_state"] == "none" and not saved["teacher_permissions"] and not saved["student_permissions"]
    duplicate = await post(client, "/preview", body)
    assert duplicate["duplicates"][0]["id"] == saved["exam_id"]
    another = await post(client, "", {**body, "preview_hash": duplicate["preview_hash"]})
    assert another["id"] != saved["id"] and another["exam_id"] != saved["exam_id"]
    assert await test_db.scalar(select(func.count()).select_from(ExamOperation)) == 2


@pytest.mark.asyncio
async def test_invalid_fields_student_state_and_cross_school_are_atomic(client, test_db, actors, exam_people):
    body = draft(exam_people)
    for invalid in ["", " "]:
        await post(client, "/preview", {**body, "name": invalid}, 422)
    for score in [0, -1, "NaN", "Infinity", "1.234"]:
        invalid = {**body, "subjects": [{**body["subjects"][0], "full_score": score}]}
        await post(client, "/preview", invalid, 422)
    await post(client, "/preview", {**body, "subjects": [body["subjects"][0], {**body["subjects"][1], "name": "数学"}]}, 422)
    (await test_db.get(Student, exam_people[1].id)).status = "departed"
    await test_db.commit()
    await post(client, "/preview", body, 422)
    foreign = Student(school_id=actors["other_school"].id, name="跨校学生")
    test_db.add(foreign); await test_db.commit()
    await post(client, "/preview", draft([foreign], count=1), 422)
    assert await test_db.scalar(select(func.count()).select_from(ExamAdministration)) == 0


@pytest.mark.asyncio
async def test_score_full_roster_template_validation_original_and_questions(client, test_db, actors, exam_people):
    exam = await create(client, exam_people, subjects=1)
    body, _ = await score_file(client, exam)
    book = load_workbook(io.BytesIO(base64.b64decode(body["content_base64"])))
    sheet = book.active
    sheet.delete_rows(3)
    output = io.BytesIO(); book.save(output)
    invalid = await post(client, f"/{exam['id']}/score-files", {**body, "content_base64": base64.b64encode(output.getvalue()).decode()})
    assert invalid["status"] == "invalid" and any("完整名单" in error for error in invalid["errors"])
    assert await test_db.scalar(select(func.count()).select_from(StudentSubjectScore)) == 0
    await post(client, f"/{exam['id']}/score-files/{invalid['id']}/confirm", status=422)
    exam, batch, original = await import_score(client, exam, questions=True)
    assert exam["open_state"] == "none"
    response = await client.get(ROOT+f"/{exam['id']}/score-files/{batch['id']}/original")
    assert response.content == original
    assert await test_db.scalar(select(func.count()).select_from(QuestionScore)) == 4
    assert all(q.full_score is None and q.lost_score is None for q in (await test_db.scalars(select(QuestionScore))).all())
    await post(client, f"/{exam['id']}/score-files", {**body, "expected_revision": exam["revision"]}, 409)


@pytest.mark.asyncio
@pytest.mark.parametrize("bad", ["formula", "negative", "overflow", "duplicate", "name", "columns", "questions"])
async def test_bad_score_files_do_not_write_scores(client, test_db, actors, exam_people, bad):
    exam = await create(client, exam_people, subjects=1)
    body, original = await score_file(client, exam)
    book = load_workbook(io.BytesIO(original)); sheet = book.active
    if bad == "formula": sheet.cell(2, 7, "=10+20")
    if bad == "negative": sheet.cell(2, 7, -1)
    if bad == "overflow": sheet.cell(2, 7, 151)
    if bad == "duplicate": sheet.cell(3, 6, sheet.cell(2, 6).value)
    if bad == "name": sheet.cell(2, 3, "错误姓名")
    if bad == "columns": sheet.cell(1, 6, "已改名")
    if bad == "questions": sheet.cell(1, 10, "1"); sheet.cell(1, 11, "1")
    output = io.BytesIO(); book.save(output); book.close()
    result = await post(client, f"/{exam['id']}/score-files", {**body, "content_base64": base64.b64encode(output.getvalue()).decode()}, 422 if bad == "formula" else 200)
    if bad != "formula": assert result["status"] == "invalid"
    assert await test_db.scalar(select(func.count()).select_from(StudentSubjectScore)) == 0


@pytest.mark.asyncio
async def test_permissions_enforce_student_subject_total_tools_and_teacher_independence(client, test_db, actors, exam_people):
    exam = await create(client, exam_people)
    exam, _, _ = await import_score(client, exam)
    exam, _, _ = await import_score(client, exam, 1)
    switch(actors, test_db, "student")
    assert (await client.get(ROOT+f"/my-exams/{exam['id']}")).status_code == 404
    dashboard = await client.get("/api/v1/platform/student/dashboard")
    assert dashboard.status_code == 200 and dashboard.json()["latest_exam"] is None
    switch(actors, test_db, "admin")
    first, second = [s["subject_id"] for s in exam["subjects"]]
    student_body = {"student_ids": [str(exam_people[0].id)], "subject_ids": [first], "mode": "open"}
    await post(client, f"/{exam['id']}/permissions/student", {**student_body, "expected_revision": exam["revision"], "preview_hash": "wrong"}, 409)
    exam = await permission(client, exam, "student", student_body)
    assert exam["open_state"] == "part" and not exam["teacher_permissions"]
    switch(actors, test_db, "student")
    results = (await client.get(ROOT+f"/my-exams/{exam['id']}")).json()
    assert len(results["scores"]) == 1 and results["scores"][0]["subject"] == "数学"
    dashboard = (await client.get("/api/v1/platform/student/dashboard")).json()
    assert dashboard["latest_exam"] is None  # Closed Chinese score must not leak through a total.
    context = ToolContext(request_id=str(uuid4()), agent_run_id=str(uuid4()), user_id=actors["student"].user_id,
        school_id=actors["school"].id, student_id=exam_people[0].id, entitlement_level="BASIC")
    total = await GetExamSummaryTool(test_db).execute(context, {"exam_id": exam["id"]})
    assert not total.ok
    subjects = await GetSubjectScoresTool(test_db).execute(context, {"exam_id": exam["id"]})
    assert subjects.ok and len(subjects.data["subjects"]) == 1
    switch(actors, test_db, "admin")
    exam = await permission(client, exam, "student", {**student_body, "subject_ids": [second]})
    switch(actors, test_db, "student")
    dashboard = (await client.get("/api/v1/platform/student/dashboard")).json()
    assert dashboard["latest_exam"]["total_score"] == 180 and len(dashboard["subjects"]) == 2
    switch(actors, test_db, "admin")
    teacher = {"teacher_id": str(actors["teacher"].id), "score_types": ["student", "classSubject"], "score_subjects": [first], "range_type": "all"}
    exam = await permission(client, exam, "teacher", {"teachers": [teacher]})
    switch(actors, test_db, "teacher")
    results = (await client.get(ROOT+f"/my-exams/{exam['id']}")).json()
    assert len(results["scores"]) == 2 and {s["subject"] for s in results["scores"]} == {"数学"}
    await post(client, f"/{exam['id']}/permissions/student", {**student_body, "expected_revision": exam["revision"]}, 403)
    await post(client, "/preview", draft(exam_people), 403)


@pytest.mark.asyncio
async def test_adjust_freezes_history_retains_data_and_does_not_expand_grants(client, test_db, actors, exam_people):
    exam = await create(client, exam_people, subjects=1)
    exam, _, _ = await import_score(client, exam)
    subject = exam["subjects"][0]["subject_id"]
    exam = await permission(client, exam, "teacher", {"teachers": [{"teacher_id": str(actors["teacher"].id), "score_types": ["student"], "score_subjects": [subject], "range_type": "all"}]})
    before_grants = await test_db.scalar(select(func.count()).select_from(ExamQueryGrant))
    year_record = await test_db.scalar(select(StudentYearClass).where(StudentYearClass.student_id == exam_people[0].id))
    year_record.class_number = "2班"; (await test_db.get(Student, exam_people[0].id)).status = "departed"
    await test_db.commit()
    body = {key: exam[key] for key in ["name", "academic_year", "exam_date", "subjects"]}
    body = {**body, "subjects": [{**exam["subjects"][0], "student_ids": [str(s.id) for s in exam_people]}], "expected_revision": exam["revision"]}
    preview = await post(client, f"/{exam['id']}/preview", body)
    assert preview["linked_data"]["score_records"] == 2
    saved = await client.put(ROOT+f"/{exam['id']}", json={**body, "preview_hash": preview["preview_hash"]})
    assert saved.status_code == 200, saved.text
    result = saved.json()
    assert result["snapshots"][str(exam_people[0].id)]["class_number"] == "1班"
    assert result["participant_count"] == 3
    assert await test_db.scalar(select(func.count()).select_from(StudentSubjectScore)) == 2
    assert await test_db.scalar(select(func.count()).select_from(ExamQueryGrant)) == before_grants
    grants = (await test_db.scalars(select(ExamQueryGrant))).all()
    assert all(g.student_id != exam_people[2].id for g in grants)
    assert (await client.put(ROOT+f"/{exam['id']}", json={**body, "preview_hash": preview["preview_hash"]})).status_code == 409


@pytest.mark.asyncio
async def test_delete_reimport_revokes_score_access_preserves_reports(client, test_db, actors, exam_people):
    exam = await create(client, exam_people, subjects=1)
    exam, batch, _ = await import_score(client, exam, questions=True)
    sub = exam["subjects"][0]["subject_id"]
    report = DiagnosisReport(school_id=actors["school"].id, student_id=exam_people[0].id, exam_id=UUID(exam["id"]), subject_id=UUID(sub),
        report_type="diagnosis", version="v1", generated_at=datetime.now(timezone.utc), structured_json={"summary": "合成报告"})
    test_db.add(report); await test_db.commit()
    exam = await permission(client, exam, "student", {"student_ids": [str(exam_people[0].id)], "subject_ids": [sub], "mode": "open"})
    exam = await permission(client, exam, "teacher", {"teachers": [{"teacher_id": str(actors["teacher"].id), "score_types": ["student"],
        "score_subjects": [sub], "diagnosis_types": ["student"], "diagnosis_subjects": [sub], "range_type": "all"}]})
    exam = await post(client, f"/{exam['id']}/score-files/{sub}/delete", {"expected_revision": exam["revision"]})
    assert not exam["score_files"] and not exam["student_permissions"]
    assert exam["teacher_permissions"][str(actors["teacher"].id)]["score_subjects"] == []
    assert exam["teacher_permissions"][str(actors["teacher"].id)]["diagnosis_subjects"] == [sub]
    assert await test_db.scalar(select(func.count()).select_from(QuestionScore)) == 0
    assert await test_db.scalar(select(func.count()).select_from(StudentExamScore)) == 0
    assert await test_db.get(DiagnosisReport, report.id) is not None
    assert (await client.get(ROOT+f"/{exam['id']}/score-files/{batch['id']}/original")).status_code == 404
    exam, _, _ = await import_score(client, exam)
    assert not exam["student_permissions"]
    assert exam["teacher_permissions"][str(actors["teacher"].id)]["score_types"] == []


@pytest.mark.asyncio
async def test_permission_excel_round_trip_and_stale_preview(client, test_db, actors, exam_people):
    exam = await create(client, exam_people, subjects=1)
    exam, _, _ = await import_score(client, exam)
    sub = exam["subjects"][0]["subject_id"]
    payload = {"expected_revision": exam["revision"], "teachers": [{"teacher_id": str(actors["teacher"].id), "score_types": ["student"], "score_subjects": [sub], "range_type": "all"}]}
    response = await client.post(ROOT+f"/{exam['id']}/permissions/teacher/preview", json=payload)
    assert response.status_code == 200
    checked = await post(client, f"/{exam['id']}/permissions/teacher/upload", {"filename": "授权.xlsx", "content_base64": base64.b64encode(response.content).decode()})
    assert checked["items"][0]["teacher_id"] == str(actors["teacher"].id) and checked["saved"] is False
    assert await test_db.scalar(select(func.count()).select_from(ExamQueryGrant)) == 0
    payload["teachers"][0]["score_types"] = ["classSubject"]
    await post(client, f"/{exam['id']}/permissions/teacher", {**payload, "preview_hash": response.headers["X-Exam-Preview-Hash"]}, 409)
    switch(actors, test_db, "foreign")
    assert (await client.get(ROOT+f"/{exam['id']}")).status_code == 404
    switch(actors, test_db, "qa")
    await post(client, "/preview", draft(exam_people), 403)


@pytest.mark.asyncio
async def test_permissions_save_directly_without_download_and_still_validate_revision_and_subjects(client, test_db, actors, exam_people):
    exam = await create(client, exam_people)
    exam, _, _ = await import_score(client, exam)
    subject, unimported = [subject["subject_id"] for subject in exam["subjects"]]
    teacher_body = {"expected_revision": exam["revision"], "teachers": [{"teacher_id": str(actors["teacher"].id),
        "score_types": ["student"], "score_subjects": [subject], "range_type": "all"}]}
    invalid = {**teacher_body, "teachers": [{**teacher_body["teachers"][0], "score_subjects": [unimported]}]}
    await post(client, f"/{exam['id']}/permissions/teacher", invalid, 422)
    assert await test_db.scalar(select(func.count()).select_from(ExamQueryGrant)) == 0
    saved = await post(client, f"/{exam['id']}/permissions/teacher", teacher_body)
    assert saved["teacher_permissions"][str(actors["teacher"].id)]["score_subjects"] == [subject]
    await post(client, f"/{exam['id']}/permissions/teacher", teacher_body, 409)
    student_body = {"expected_revision": saved["revision"], "student_ids": [str(exam_people[0].id)], "subject_ids": [subject], "mode": "open"}
    saved = await post(client, f"/{exam['id']}/permissions/student", student_body)
    assert saved["student_permissions"][str(exam_people[0].id)]["subjects"] == [subject]
    await post(client, f"/{exam['id']}/permissions/student", student_body, 409)
    saved = await post(client, f"/{exam['id']}/permissions/student", {**student_body, "expected_revision": saved["revision"], "mode": "close"})
    assert not saved["student_permissions"]


@pytest.mark.asyncio
async def test_edit_exam_metadata_and_year_rebuilds_year_placements_and_preserves_linked_data(client, test_db, actors, exam_people):
    exam = await create(client, exam_people, subjects=1)
    exam, _, _ = await import_score(client, exam)
    subject = exam["subjects"][0]["subject_id"]
    exam = await permission(client, exam, "teacher", {"teachers": [{"teacher_id": str(actors["teacher"].id),
        "score_types": ["student"], "score_subjects": [subject], "range_type": "all"}]})
    grants_before = await test_db.scalar(select(func.count()).select_from(ExamQueryGrant))
    test_db.add(StudentYearClass(school_id=actors["school"].id, student_id=exam_people[0].id,
        academic_year="2025-2026", grade="初中三年级", class_number="8班"))
    await test_db.commit()
    body = {key: exam[key] for key in ["name", "academic_year", "exam_date", "subjects"]}
    body.update(name="修改后的考试", academic_year="2025-2026", exam_date="2025-12-18", expected_revision=exam["revision"])
    await post(client, f"/{exam['id']}/preview", body, 422)
    body["subjects"] = [{**body["subjects"][0], "student_ids": [str(exam_people[0].id)]}]
    checked = await post(client, f"/{exam['id']}/preview", body)
    response = await client.put(ROOT+f"/{exam['id']}", json={**body, "preview_hash": checked["preview_hash"]})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["name"] == "修改后的考试" and result["academic_year"] == "2025-2026" and result["exam_date"] == "2025-12-18"
    assert result["snapshots"][str(exam_people[0].id)]["class_number"] == "8班"
    assert result["snapshots"][str(exam_people[0].id)]["academic_year"] == "2025-2026"
    assert str(exam_people[1].id) not in result["snapshots"]
    stored = await test_db.get(Exam, UUID(exam["id"]))
    assert stored.academic_year == "2025-2026" and stored.start_date.isoformat() == "2025-12-18" and stored.name == "修改后的考试"
    placement = await test_db.scalar(select(StudentExamPlacement).where(StudentExamPlacement.exam_id == stored.id,
        StudentExamPlacement.student_id == exam_people[0].id))
    assert placement.academic_year == "2025-2026" and placement.grade == "初中三年级" and placement.class_number == "8班"
    assert await test_db.scalar(select(func.count()).select_from(StudentSubjectScore)) == 2
    assert await test_db.scalar(select(func.count()).select_from(ExamQueryGrant)) == grants_before
    listing = await client.get(ROOT+"?academic_year=2025-2026&search=修改后的考试")
    assert listing.json()["total"] == 1


@pytest.mark.asyncio
async def test_teacher_default_options_use_profile_roles_and_raw_scope_without_class_references(client, test_db, actors, exam_people):
    test_db.add(TeacherProfile(id=actors["teacher"].id, school_id=actors["school"].id,
        fields={"任课年级班级": "数学:10.1", "班主任年级班级": "10.1"}, duties=["TEACHER", "HOMEROOM_TEACHER"]))
    await test_db.commit()
    exam = await create(client, exam_people)
    exam, _, _ = await import_score(client, exam)
    response = await client.get(ROOT+f"/teachers?exam_id={exam['id']}")
    assert response.status_code == 200, response.text
    data = response.json()
    teacher = next(teacher for teacher in data["items"] if teacher["id"] == str(actors["teacher"].id))
    assert teacher["role_labels"] == ["任课教师", "班主任"]
    assert teacher["recommendation"]["score_types"] == ["student", "classSubject", "classAll"]
    assert teacher["recommendation"]["score_subjects"] == [exam["subjects"][0]["subject_id"]]
    assert teacher["recommendation"]["diagnosis_subjects"] == [subject["subject_id"] for subject in exam["subjects"]]
    assert teacher["recommendation"]["classes"] == ["高中一年级 / 1班"]
    assert teacher["filter_scopes"] == [
        {"role": "TEACHER", "grade": "高中一年级", "class_number": "1", "subject": "数学"},
        {"role": "HOMEROOM_TEACHER", "grade": "高中一年级", "class_number": "1", "subject": ""},
    ]
    assert "GENERAL_DIRECTOR" in {rule["role"] for rule in data["default_rules"]}
    switch(actors, test_db, "foreign")
    assert (await client.get(ROOT+f"/teachers?exam_id={exam['id']}")).status_code == 404


@pytest.mark.asyncio
async def test_default_rules_download_contains_complete_rules_without_saving_permissions(client, test_db, actors):
    response = await client.get(ROOT+"/teacher-default-rules")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert response.headers["content-disposition"].startswith("attachment;")
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["默认选项规则", "使用说明"]
    sheet = book["默认选项规则"]
    assert list(next(sheet.values)) == ["教师角色", "默认成绩类型", "默认诊断类型", "默认科目", "默认学生范围"]
    rows = {row[0]: row for row in list(sheet.values)[1:]}
    assert len(rows) == 8 and "总务主任" in rows
    assert "班级综合成绩" in rows["班主任"][1] and "班级综合诊断" in rows["班主任"][2]
    assert sheet["B2"].alignment.wrap_text
    assert book["使用说明"].max_row > 5
    book.close()
    assert await test_db.scalar(select(func.count()).select_from(ExamQueryGrant)) == 0
    assert await test_db.scalar(select(func.count()).select_from(ExamOperation)) == 0
    switch(actors, test_db, "teacher")
    assert (await client.get(ROOT+"/teacher-default-rules")).status_code == 403


@pytest.mark.asyncio
async def test_report_original_access_follows_exam_authorization_without_legacy_entitlement(client, test_db, actors, exam_people):
    exam = await create(client, exam_people)
    exam, _, _ = await import_score(client, exam)
    sub = exam["subjects"][0]["subject_id"]
    content = b"%PDF-1.4\n% synthetic immutable original\n%%EOF"
    asset = PrivateAsset(school_id=actors["school"].id, filename="合成诊断.pdf", media_type="application/pdf", byte_size=len(content),
        content_hash=hashlib.sha256(content).hexdigest(), storage_key="exam-qa-"+uuid4().hex, content=content, created_by=actors["admin"].id)
    test_db.add(asset); await test_db.flush()
    report = DiagnosisReport(school_id=actors["school"].id, student_id=exam_people[0].id, exam_id=UUID(exam["id"]), subject_id=UUID(sub),
        report_type="diagnosis", version="v1", generated_at=datetime.now(timezone.utc), structured_json={"summary":"合成诊断"})
    test_db.add(report); await test_db.flush()
    report_id, asset_id = report.id, asset.id
    test_db.add(ReportAttachment(school_id=actors["school"].id, report_id=report_id, asset_id=asset_id))
    await test_db.commit()
    exam = await permission(client, exam, "student", {"student_ids":[str(exam_people[0].id)],"subject_ids":[sub],"mode":"open"})
    switch(actors, test_db, "student")
    response = await client.get(ROOT+f"/my-exams/{exam['id']}/reports/{report_id}/original")
    assert response.status_code == 200 and response.content == content
    assert (await client.get("/api/v1/platform/student/diagnosis")).status_code == 200
    switch(actors, test_db, "admin")
    exam = await permission(client, exam, "teacher", {"teachers":[{"teacher_id":str(actors["teacher"].id),"diagnosis_types":["student"],"diagnosis_subjects":[sub],"range_type":"all"}]})
    switch(actors, test_db, "teacher")
    result = await client.get(ROOT+f"/my-exams/{exam['id']}")
    assert result.json()["scores"] == [] and len(result.json()["reports"]) == 1
    assert (await client.get(ROOT+f"/my-exams/{exam['id']}/reports/{report_id}/original")).content == content
    switch(actors, test_db, "admin")
    exam = await post(client, f"/{exam['id']}/permissions/teacher/{actors['teacher'].id}/revoke", {"expected_revision":exam["revision"]})
    exam = await permission(client, exam, "student", {"student_ids":[str(exam_people[0].id)],"subject_ids":[sub],"mode":"close"})
    switch(actors, test_db, "student")
    assert (await client.get(ROOT+f"/my-exams/{exam['id']}/reports/{report_id}/original")).status_code == 404
    assert (await client.get(f"/api/v1/platform/education/assets/{asset_id}")).status_code in {403,404}
    switch(actors, test_db, "teacher")
    assert (await client.get(ROOT+f"/my-exams/{exam['id']}/reports/{report_id}/original")).status_code == 404


@pytest.mark.asyncio
async def test_teacher_scope_prevents_totals_and_other_subjects_in_both_personal_and_combined_views(client, test_db, actors, exam_people):
    exam = await create(client, exam_people)
    exam, _, _ = await import_score(client, exam)
    exam, _, _ = await import_score(client, exam, 1)
    first, second = [s["subject_id"] for s in exam["subjects"]]
    p = {"teacher_id":str(actors["teacher"].id),"score_types":["student","classAll"],"score_subjects":[first],"range_type":"all"}
    exam = await permission(client, exam, "teacher", {"teachers":[p]})
    switch(actors, test_db, "teacher")
    personal = (await client.get(ROOT+f"/my-exams/{exam['id']}")).json()
    assert len(personal["scores"]) == 2 and personal["totals"] == []
    combined = (await client.get(ROOT+f"/my-exams/{exam['id']}?view_type=classAll")).json()
    assert combined["totals"] == []
    # The shared session guard protects forgotten/count-only queries too.
    assert await test_db.scalar(select(func.count(StudentSubjectScore.id))) == 2
    assert await test_db.scalar(select(func.count(StudentExamScore.id))) == 0
    switch(actors, test_db, "admin")
    exam = await permission(client, exam, "teacher", {"teachers":[{**p,"score_subjects":[first,second]}]})
    switch(actors, test_db, "teacher")
    combined = (await client.get(ROOT+f"/my-exams/{exam['id']}?view_type=classAll")).json()
    assert len(combined["totals"]) == 2
    assert {float(t["total_score"]) for t in combined["totals"]} == {180,190}


@pytest.mark.asyncio
async def test_legacy_editor_cannot_bypass_managed_exam_preview(client, actors, exam_people):
    exam = await create(client, exam_people, subjects=1)
    sub = exam["subjects"][0]["subject_id"]
    response = await client.put(f"/api/v1/platform/education/exams/{exam['id']}/subjects", json={"subject_id":sub,"full_score":100})
    assert response.status_code == 409 and response.json()["error"]["code"] == "EXAM_WORKFLOW_REQUIRED"


@pytest.mark.asyncio
async def test_whole_exam_report_stays_hidden_until_every_participating_subject_is_open(client, test_db, actors, exam_people):
    exam = await create(client, exam_people)
    exam, _, _ = await import_score(client, exam)
    report = DiagnosisReport(school_id=actors["school"].id, student_id=exam_people[0].id, exam_id=UUID(exam["id"]),
        subject_id=None, report_type="diagnosis", version="all", generated_at=datetime.now(timezone.utc), structured_json={"summary":"全科合成报告"})
    test_db.add(report); await test_db.commit()
    report_id = str(report.id)
    first, second = [s["subject_id"] for s in exam["subjects"]]
    body = {"student_ids":[str(exam_people[0].id)],"subject_ids":[first],"mode":"open"}
    exam = await permission(client, exam, "student", body)
    switch(actors, test_db, "student")
    assert (await client.get(ROOT+f"/my-exams/{exam['id']}")).json()["reports"] == []
    assert (await client.get("/api/v1/platform/student/diagnosis")).json() == []
    switch(actors, test_db, "admin")
    exam, _, _ = await import_score(client, exam, 1)
    exam = await permission(client, exam, "student", {**body,"subject_ids":[second]})
    switch(actors, test_db, "student")
    assert (await client.get(ROOT+f"/my-exams/{exam['id']}")).json()["reports"][0]["id"] == report_id


@pytest.mark.asyncio
async def test_teacher_authorization_rejects_subjects_with_no_selected_participants_atomically(client, test_db, actors, exam_people):
    placement = await test_db.scalar(select(StudentYearClass).where(StudentYearClass.student_id == exam_people[1].id))
    placement.class_number = "2班"
    await test_db.commit()
    body = draft(exam_people)
    body["subjects"][0]["student_ids"] = [str(exam_people[0].id)]
    body["subjects"][1]["student_ids"] = [str(exam_people[1].id)]
    checked = await post(client, "/preview", body)
    exam = await post(client, "", {**body, "preview_hash": checked["preview_hash"]})
    exam, _, _ = await import_score(client, exam)
    invalid = {"expected_revision": exam["revision"], "teachers": [{"teacher_id": str(actors["teacher"].id),
        "score_types": ["student"], "score_subjects": [exam["subjects"][0]["subject_id"]],
        "range_type": "class", "classes": ["高中一年级 / 2班"]}]}
    response = await client.post(ROOT+f"/{exam['id']}/permissions/teacher", json=invalid)
    assert response.status_code == 422 and "没有参试学生" in response.text
    assert await test_db.scalar(select(func.count()).select_from(ExamQueryGrant)) == 0
    current = (await client.get(ROOT+f"/{exam['id']}")).json()
    assert current["revision"] == exam["revision"] and not current["teacher_permissions"]
    invalid["teachers"][0]["classes"] = ["高中一年级 / 1班"]
    saved = await post(client, f"/{exam['id']}/permissions/teacher", invalid)
    assert saved["teacher_permissions"][str(actors["teacher"].id)]["frozen_student_ids"] == [str(exam_people[0].id)]


@pytest.mark.asyncio
@pytest.mark.parametrize("kind", ["teacher", "student"])
async def test_authorization_upload_rejects_name_and_person_id_mismatches(client, test_db, actors, exam_people, kind):
    exam = await create(client, exam_people)
    exam, _, _ = await import_score(client, exam)
    subject = exam["subjects"][0]["subject_id"]
    body = {"expected_revision": exam["revision"]}
    if kind == "teacher":
        body["teachers"] = [{"teacher_id": str(actors["teacher"].id), "score_types": ["student"], "score_subjects": [subject], "range_type": "all"}]
    else:
        body.update(student_ids=[str(exam_people[0].id)], subject_ids=[subject], mode="open")
    response = await client.post(ROOT+f"/{exam['id']}/permissions/{kind}/preview", json=body)
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content))
    sheet = book.active
    header_row = next(row[0].row for row in sheet if any(cell.value == "姓名" for cell in row))
    sheet.cell(header_row+1, 3, "另一个人的姓名")
    output = io.BytesIO(); book.save(output); book.close()
    result = await post(client, f"/{exam['id']}/permissions/{kind}/upload", {
        "filename": "错误姓名授权.xlsx", "content_base64": base64.b64encode(output.getvalue()).decode()}, 422)
    assert "姓名与人员ID不匹配" in str(result)
    assert await test_db.scalar(select(func.count()).select_from(ExamQueryGrant)) == 0
