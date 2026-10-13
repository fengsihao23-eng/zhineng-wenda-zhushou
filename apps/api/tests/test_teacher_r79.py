"""Acceptance for the supplied r79 rules and v3 teacher archive prototype."""
import base64
import io
from uuid import UUID, uuid4
import pytest
from openpyxl import load_workbook
from sqlalchemy import select, func
from sqlalchemy.exc import IntegrityError
from app.db.models import User, TeacherProfile, TeacherDeletion, TeacherChangeTask, RosterImport, TeachingAssignment
from app.core.security import get_password_hash, verify_password
from app.api.deps import get_current_user
from app.main import app
from app.services.roster_excel import TEACHER_COLUMNS, TEACHER_CHANGE_COLUMNS, workbook_bytes, duty_codes
from test_education_workbench import actors, post, P  # noqa: F401
from test_roster_workflows import teacher, upload, confirm, classroom


async def delete_teachers(client, identifiers, reason="change", **values):
    return await post(client, "/roster/teachers/delete", {"teacher_ids": identifiers, "reason": reason, "confirmed": True, **values})


async def upload_task(client, task_id, rows, columns=None):
    headers = columns or TEACHER_CHANGE_COLUMNS
    content = workbook_bytes(headers, [[row.get(column, "") for column in headers] for row in rows])
    return await post(client, f"/roster/teachers/change-tasks/{task_id}/imports", {
        "filename": "教师更新.xlsx", "content_base64": base64.b64encode(content).decode(),
    })


async def task_rows(client, identifier):
    task = (await client.get(P + f"/roster/teachers/change-tasks/{identifier}")).json()
    return [{**row["fields"], "教师ID": row["platform_id"]} for row in task["snapshots"]]


async def confirm_task(client, batch, status=200, **values):
    return await confirm(client, batch, status=status, identity_confirmed=True, changes_confirmed=True, **values)


@pytest.mark.asyncio
async def test_teacher_phone_contract_exact_text_duplicates_and_original_download(client, test_db, actors):
    bad = await upload(client, [teacher("account-not-phone"), teacher("139****0001"), teacher("13900000001", 手机号码="139****0001")])
    assert bad["status"] == "invalid"
    assert len(bad["report"]["issues"]) == 2
    assert {issue["fields"] for issue in bad["report"]["issues"]} == {"教师账号"}
    await confirm(client, bad, status=409)
    assert await test_db.scalar(select(User).where(User.username == "13900000001")) is None
    old = await confirm(client, await upload(client, [teacher("13900000001", "同名", 任课年级班级="数学:10.1,10.2;", 手机号码="020-12345678")]))
    candidate = await upload(client, [teacher("13900000002", "同名", 任课年级班级="数学:10.1,10.2;"), teacher("13900000003", "同名", 任课年级班级="数学:10.2,10.1;"), teacher("13900000004", "同名", 任课年级班级="数学:10.3;"), teacher("13900000005", "同名"), teacher("13900000006", " 同名 ", 任课年级班级="数学:10.1,10.2;")])
    assert [row["row_number"] for row in candidate["report"]["suspicious"]] == [2]
    await confirm(client, candidate, status=422)
    done = await confirm(client, candidate, duplicate_decisions={2: False})
    assert done["report"]["success_count"] == 4 and done["report"]["skipped_rows"] == [2]
    original = await client.get(P + f"/roster/imports/{old['id']}/original")
    book = load_workbook(io.BytesIO(original.content))
    assert [cell.value for cell in book.active[1]] == TEACHER_COLUMNS
    assert book.active["M2"].value == "020-12345678"
    book.close()
    found = (await client.get(P + "/roster/teachers/records?search=020-12345678")).json()
    assert found["total"] == 1
    receipts = (await client.get(P + "/roster/teachers/imports?paginated=true&search=13900000001")).json()
    assert receipts["total"] == 2 and old["id"] in {row["id"] for row in receipts["items"]}


@pytest.mark.asyncio
async def test_masked_contacts_survive_import_and_original_download(client, test_db, actors):
    contacts = ["139****0001", "139＊＊＊＊0002", "139••••0003", "139●●●●0004", "", "13900000006"]
    batch = await upload(client, [teacher(f"1390000000{index}", f"教师{index}", 手机号码=contact)
                                 for index, contact in enumerate(contacts, 1)])
    assert batch["status"] == "ready" and batch["report"]["issues"] == []
    done = await confirm(client, batch)
    assert done["report"]["success_count"] == len(contacts)
    for imported, contact in zip(done["report"]["imported"], contacts):
        identifier = UUID(imported["id"])
        assert (await test_db.get(TeacherProfile, identifier)).fields["手机号码"] == contact
        assert (await test_db.get(User, identifier)).phone == (contact or None)
    original = await client.get(P + f"/roster/imports/{batch['id']}/original")
    book = load_workbook(io.BytesIO(original.content))
    assert [book.active.cell(index, 13).value or "" for index in range(2, 8)] == contacts
    book.close()


@pytest.mark.asyncio
async def test_masked_contact_change_retains_teacher_identity_and_versions(client, test_db, actors):
    original = await confirm(client, await upload(client, [teacher("13900000001", 手机号码="139****0001")]))
    identifier = original["report"]["imported"][0]["id"]
    deletion = await delete_teachers(client, [identifier])
    template = await client.get(P + f"/roster/teachers/change-tasks/{deletion['task_id']}/template")
    book = load_workbook(io.BytesIO(template.content))
    assert book.active["M2"].value == "139****0001"
    book.close()
    rows = await task_rows(client, deletion["task_id"])
    rows[0]["手机号码"] = "138****0002"
    batch = await upload_task(client, deletion["task_id"], rows)
    assert batch["status"] == "ready" and batch["report"]["issues"] == []
    done = await confirm_task(client, batch)
    assert done["report"]["imported"][0]["id"] == identifier
    user = await test_db.get(User, UUID(identifier))
    assert user.phone == "138****0002" and user.username == "13900000001"
    assert (await test_db.get(TeacherProfile, user.id)).fields["手机号码"] == "138****0002"
    snapshot = await test_db.get(TeacherDeletion, UUID(deletion["snapshots"][0]["id"]))
    assert snapshot.fields["手机号码"] == "139****0001"
    assert snapshot.updated_fields["手机号码"] == "138****0002"


@pytest.mark.asyncio
async def test_teacher_template_download_works_without_private_workbook(client, actors, monkeypatch, tmp_path):
    from app.api.v1.endpoints import roster as endpoints
    monkeypatch.setattr(endpoints, "TEMPLATES", tmp_path)
    response = await client.get(P + "/roster/teachers/template")
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    assert [cell.value for cell in book.active[1]] == TEACHER_COLUMNS
    assert book.active.max_row == 1  # No real staff data is needed to deploy or download.
    book.close()
    batch = await post(client, "/roster/teachers/imports", {"filename": "教师资料模板.xlsx", "content_base64": base64.b64encode(response.content).decode()})
    assert batch["status"] == "invalid" and batch["report"]["total_rows"] == 0


def test_teacher_description_never_grants_roles_or_management():
    assert duty_codes({"说明": "学校管理员;考试管理员;班主任;备课组长;年级长;校长;教务主任"}) == ["TEACHER"]
    assert "HOMEROOM_TEACHER" in duty_codes({"班主任年级班级": "10.1"})
    assert "PRINCIPAL" in duty_codes({"是否校长": "是"})


@pytest.mark.asyncio
@pytest.mark.parametrize("reason", ["depart", "disable", "incorrect", "other"])
async def test_pure_deletion_does_not_bind_future_import_or_reuse_identity(client, test_db, actors, reason):
    done = await confirm(client, await upload(client, [teacher("13900000001")]))
    original_id = done["report"]["imported"][0]["id"]
    deletion = await delete_teachers(client, [original_id], reason, reason_note="重复建档" if reason == "other" else "")
    assert deletion["task_id"] is None
    assert await test_db.scalar(select(func.count()).select_from(TeacherChangeTask)) == 0
    assert (await test_db.get(User, UUID(original_id))).status == "deleted"
    fresh = await confirm(client, await upload(client, [teacher("13900000001")]))
    assert fresh["report"]["imported"][0]["id"] != original_id
    snapshot = (await client.get(P + f"/roster/teachers/deletions/{deletion['snapshots'][0]['id']}")).json()
    assert snapshot["reimported_user_id"] is None and snapshot["reason"] == reason
    assert snapshot["fields"]["教师账号"] == "13900000001"


@pytest.mark.asyncio
async def test_delete_requires_reason_verification_and_is_atomic_and_idempotent(client, test_db, actors):
    done = await confirm(client, await upload(client, [teacher("13900000001"), teacher("13900000002", "乙")]))
    identifiers = [row["id"] for row in done["report"]["imported"]]
    await post(client, "/roster/teachers/delete", {"teacher_ids": identifiers, "confirmed": True}, status=422)
    await post(client, "/roster/teachers/delete", {"teacher_ids": identifiers, "reason": "other", "reason_note": " ", "confirmed": True}, status=422)
    await post(client, "/roster/teachers/delete", {"teacher_ids": [identifiers[0], str(uuid4())], "reason": "depart", "confirmed": True}, status=404)
    assert (await test_db.get(User, UUID(identifiers[0]))).status == "active"
    assert await test_db.scalar(select(func.count()).select_from(TeacherDeletion)) == 0
    await post(client, f"/roster/teachers/records/{identifiers[0]}/actions", {"action": "delete", "confirmed": True}, status=422)
    key = uuid4()
    body = {"teacher_ids": identifiers, "reason": "change", "confirmed": True}
    first = await post(client, "/roster/teachers/delete", body, key=key)
    replay = await post(client, "/roster/teachers/delete", body, key=key)
    assert replay == first
    assert await test_db.scalar(select(func.count()).select_from(TeacherDeletion)) == 2
    assert await test_db.scalar(select(func.count()).select_from(TeacherChangeTask)) == 1


@pytest.mark.asyncio
@pytest.mark.parametrize("changed_account", [False, True])
@pytest.mark.parametrize("initial_password", [False, True])
async def test_single_task_reuses_identity_resets_password_and_preserves_versions(client, test_db, actors, changed_account, initial_password):
    await classroom(client)
    original = await confirm(client, await upload(client, [teacher("13900000001", 任课年级班级="数学:10.1;")]))
    identifier = original["report"]["imported"][0]["id"]
    user = await test_db.get(User, UUID(identifier))
    if not initial_password:
        user.password_hash, user.must_change_password = get_password_hash("previous-changed-password"), False
        await test_db.commit()
    tokens = (await client.post("/api/v1/auth/login", json={"username": user.username, "password": "previous-changed-password" if not initial_password else user.username})).json()
    original_version = user.password_version
    deletion = await delete_teachers(client, [identifier])
    template = await client.get(P + f"/roster/teachers/change-tasks/{deletion['task_id']}/template")
    book = load_workbook(io.BytesIO(template.content))
    assert [cell.value for cell in book.active[1]] == TEACHER_CHANGE_COLUMNS
    assert book.active["Y2"].value == str(user.platform_id) and book.active["Y2"].protection.locked
    assert not book.active["F2"].protection.locked
    book.close()
    rows = await task_rows(client, deletion["task_id"])
    new_account = "13900000002" if changed_account else "13900000001"
    rows[0].update({"教师姓名": "变更后的姓名", "教师账号": new_account, "任课年级班级": ""})
    batch = await upload_task(client, deletion["task_id"], rows)
    assert batch["status"] == "ready"
    await confirm(client, batch, status=422)  # Identity and changes require explicit review.
    done = await confirm_task(client, batch)
    assert done["report"]["imported"][0]["id"] == identifier
    user = await test_db.get(User, UUID(identifier))
    assert user.status == "active" and user.must_change_password
    assert user.password_version > original_version and verify_password(new_account, user.password_hash)
    snapshot = await test_db.get(TeacherDeletion, UUID(deletion["snapshots"][0]["id"]))
    assert snapshot.fields["教师姓名"] == "合成教师甲"
    assert snapshot.updated_fields["教师姓名"] == "变更后的姓名"
    assert snapshot.reimported_user_id == user.id
    assert done["report"]["versions"][0]["teacher_id"] == identifier
    assert (await test_db.scalar(select(TeachingAssignment).where(TeachingAssignment.teacher_user_id == user.id))).status == "inactive"
    replay = await confirm_task(client, batch)
    assert replay["report"] == done["report"]
    current_override = app.dependency_overrides.pop(get_current_user)
    try:
        assert (await client.get("/api/v1/auth/me", headers={"Authorization": "Bearer " + tokens["access_token"]})).status_code == 401
        assert (await client.post("/api/v1/auth/refresh", json={"refresh_token": tokens["refresh_token"]})).status_code == 401
    finally:
        app.dependency_overrides[get_current_user] = current_override
    # Successive tasks may snapshot the same stable ID without losing old versions.
    next_deletion = await delete_teachers(client, [identifier])
    assert next_deletion["task_id"] != deletion["task_id"]
    assert await test_db.scalar(select(func.count()).select_from(TeacherDeletion).where(TeacherDeletion.user_id == user.id)) == 2


@pytest.mark.asyncio
async def test_repeated_teacher_change_preserves_unchanged_teaching_history(client, test_db, actors):
    await classroom(client)
    original = await confirm(client, await upload(client, [teacher("13900000001", 任课年级班级="数学:10.1;")]))
    identifier = original["report"]["imported"][0]["id"]
    original_grant = await test_db.scalar(select(TeachingAssignment).where(TeachingAssignment.teacher_user_id == UUID(identifier)))
    original_started = original_grant.starts_at
    for version in range(1, 4):
        deletion = await delete_teachers(client, [identifier])
        rows = await task_rows(client, deletion["task_id"])
        rows[0]["教师姓名"] = f"保留任课关系教师第{version}版"
        batch = await upload_task(client, deletion["task_id"], rows)
        assert batch["status"] == "ready" and batch["report"]["issues"] == []
        done = await confirm_task(client, batch)
        assert done["report"]["imported"][0]["id"] == identifier
        grants = (await test_db.scalars(select(TeachingAssignment).where(TeachingAssignment.teacher_user_id == UUID(identifier)))).all()
        assert len(grants) == version + 1
        assert sum(grant.status == "active" for grant in grants) == 1
        assert all(grant.expires_at is not None for grant in grants if grant.status == "inactive")
        assert original_grant.status == "inactive" and original_grant.starts_at == original_started
        replay = await confirm_task(client, batch)
        assert replay["report"] == done["report"]
        assert await test_db.scalar(select(func.count()).select_from(TeachingAssignment).where(TeachingAssignment.teacher_user_id == UUID(identifier))) == version + 1
    active = next(grant for grant in grants if grant.status == "active")
    with pytest.raises(IntegrityError):
        async with test_db.begin_nested():
            test_db.add(TeachingAssignment(school_id=active.school_id, teacher_user_id=active.teacher_user_id,
                                          class_id=active.class_id, subject_id=active.subject_id, status="active"))
            await test_db.flush()


@pytest.mark.asyncio
async def test_original_downloads_remain_task_bound_when_alternating_single_and_bulk_tasks(client, actors):
    original = await confirm(client, await upload(client, [teacher("13900000001", "下载甲"), teacher("13900000002", "下载乙"), teacher("13900000003", "下载丙")]))
    identifiers = [row["id"] for row in original["report"]["imported"]]
    bulk = await delete_teachers(client, identifiers[:2])
    single = await delete_teachers(client, identifiers[2:])
    filenames = set()
    for _ in range(5):
        for deletion in (bulk, single):
            response = await client.get(P + f"/roster/teachers/change-tasks/{deletion['task_id']}/template", headers={"Origin": "http://localhost"})
            assert response.status_code == 200
            assert response.headers["x-teacher-task-id"] == deletion["task_id"]
            assert "X-Teacher-Task-ID" in response.headers["access-control-expose-headers"]
            assert response.headers["cache-control"] == "no-store"
            assert deletion["task_id"] in response.headers["content-disposition"]
            filenames.add(response.headers["content-disposition"])
            book = load_workbook(io.BytesIO(response.content))
            assert book.active.max_row == deletion["count"] + 1
            assert {(book.active.cell(row, 4).value, book.active.cell(row, 25).value)
                    for row in range(2, book.active.max_row + 1)} == {
                        (row["teacher_name"], row["platform_id"]) for row in deletion["snapshots"]}
            book.close()
    assert len(filenames) == 2


@pytest.mark.asyncio
@pytest.mark.parametrize("scenario", ["missing_row", "extra_row", "missing_id", "duplicate_id", "foreign_id", "wrong_columns"])
async def test_task_rejects_incomplete_or_modified_ids_and_retains_task(client, test_db, actors, scenario):
    done = await confirm(client, await upload(client, [teacher("13900000001"), teacher("13900000002", "乙")]))
    identifiers = [row["id"] for row in done["report"]["imported"]]
    deletion = await delete_teachers(client, identifiers)
    rows = await task_rows(client, deletion["task_id"])
    if scenario == "missing_row": rows.pop()
    if scenario == "extra_row": rows.append(dict(rows[0]))
    if scenario == "missing_id": rows[0]["教师ID"] = ""
    if scenario == "duplicate_id": rows[1]["教师ID"] = rows[0]["教师ID"]
    if scenario == "foreign_id": rows[0]["教师ID"] = str(uuid4())
    batch = await upload_task(client, deletion["task_id"], rows, TEACHER_COLUMNS if scenario == "wrong_columns" else None)
    assert batch["status"] == "invalid"
    await confirm_task(client, batch, status=409)
    task = await test_db.get(TeacherChangeTask, UUID(deletion["task_id"]))
    assert task.status == "pending"
    for identifier in identifiers:
        assert (await test_db.get(User, UUID(identifier))).status == "deleted"


@pytest.mark.asyncio
async def test_batch_task_requires_bijection_matching_ids_and_never_allows_skipping(client, test_db, actors):
    originals = await confirm(client, await upload(client, [teacher("13900000001", "甲", 任课年级班级="数学:10.1;"), teacher("13900000002", "乙")]))
    identifiers = [row["id"] for row in originals["report"]["imported"]]
    deletion = await delete_teachers(client, identifiers)
    await confirm(client, await upload(client, [teacher("13900000003", "甲", 任课年级班级="数学:10.1;")]))
    rows = await task_rows(client, deletion["task_id"])
    batch = await upload_task(client, deletion["task_id"], rows)
    suspected = {row["row_number"]: True for row in batch["report"]["suspicious"]}
    assert len(suspected) == 1
    mapping = {row["row_number"]: row["values"]["教师ID"] for row in batch["rows"]}
    await confirm_task(client, batch, status=422, teacher_replacements=mapping, duplicate_decisions={number: False for number in suspected})
    await confirm_task(client, batch, status=422, duplicate_decisions=suspected)
    await confirm_task(client, batch, status=422, teacher_replacements={2: identifiers[0], 3: identifiers[0]}, duplicate_decisions=suspected)
    swapped = {number: value for number, value in zip(mapping, reversed(list(mapping.values())))}
    await confirm_task(client, batch, status=422, teacher_replacements=swapped, duplicate_decisions=suspected)
    done = await confirm_task(client, batch, teacher_replacements=mapping, duplicate_decisions=suspected)
    assert done["report"]["success_count"] == 2 and done["report"]["skipped_count"] == 0
    assert {row["id"] for row in done["report"]["imported"]} == set(identifiers)


@pytest.mark.asyncio
async def test_task_write_failure_rolls_back_entire_batch_and_retry_succeeds(client, test_db, actors, monkeypatch):
    from app.services import roster_workbench
    from app.core.errors import ApiError
    done = await confirm(client, await upload(client, [teacher("13900000001"), teacher("13900000002", "乙")]))
    identifiers = [row["id"] for row in done["report"]["imported"]]
    deletion = await delete_teachers(client, identifiers)
    rows = await task_rows(client, deletion["task_id"])
    for row in rows: row["教师姓名"] += "已更新"
    batch = await upload_task(client, deletion["task_id"], rows)
    mapping = {row["row_number"]: row["values"]["教师ID"] for row in batch["rows"]}
    grants = roster_workbench.grant_roles
    calls = 0
    async def fail_second(db, user, roles):
        nonlocal calls
        calls += 1
        if calls == 2: raise ApiError(422, "TEST_WRITE_FAILURE", "合成第二行写入失败")
        await grants(db, user, roles)
    monkeypatch.setattr(roster_workbench, "grant_roles", fail_second)
    await confirm_task(client, batch, status=422, teacher_replacements=mapping)
    task = await test_db.get(TeacherChangeTask, UUID(deletion["task_id"]))
    assert task.status == "pending"
    for identifier in identifiers:
        assert (await test_db.get(User, UUID(identifier))).status == "deleted"
        assert "已更新" not in (await test_db.get(TeacherProfile, UUID(identifier))).fields["教师姓名"]
    assert (await test_db.get(RosterImport, UUID(batch["id"]))).status == "ready"
    monkeypatch.setattr(roster_workbench, "grant_roles", grants)
    retry = await confirm_task(client, batch, teacher_replacements=mapping)
    assert retry["report"]["success_count"] == 2


@pytest.mark.asyncio
async def test_task_confirmation_rechecks_accounts_and_retains_task_when_conflicted(client, test_db, actors):
    done = await confirm(client, await upload(client, [teacher("13900000001")]))
    deletion = await delete_teachers(client, [done["report"]["imported"][0]["id"]])
    rows = await task_rows(client, deletion["task_id"])
    batch = await upload_task(client, deletion["task_id"], rows)
    await confirm(client, await upload(client, [teacher("13900000001", "新教师")]))
    await confirm_task(client, batch, status=409)
    task = await test_db.get(TeacherChangeTask, UUID(deletion["task_id"]))
    assert task.status == "pending"
    assert (await test_db.get(RosterImport, UUID(batch["id"]))).status == "invalid"


@pytest.mark.asyncio
async def test_tasks_snapshots_originals_and_mutations_are_school_scoped(client, test_db, actors):
    done = await confirm(client, await upload(client, [teacher("13900000001")]))
    identifier = done["report"]["imported"][0]["id"]
    deletion = await delete_teachers(client, [identifier])
    rows = await task_rows(client, deletion["task_id"])
    for role in ("teacher", "student"):
        actors["switch"](role)
        assert (await client.get(P + f"/roster/teachers/change-tasks/{deletion['task_id']}")).status_code == 403
        assert (await client.get(P + f"/roster/imports/{done['id']}/original")).status_code == 403
    actors["switch"]("qa")
    await post(client, "/roster/teachers/delete", {"teacher_ids": [identifier], "reason": "depart", "confirmed": True}, status=403)
    actors["switch"]("foreign")
    assert (await client.get(P + f"/roster/teachers/change-tasks/{deletion['task_id']}")).status_code == 404
    assert (await client.get(P + f"/roster/teachers/change-tasks/{deletion['task_id']}/template")).status_code == 404
    assert (await client.get(P + f"/roster/imports/{done['id']}/original")).status_code == 404
    actors["switch"]("admin")
    new_batch = await upload(client, [teacher("13900000002")])
    await confirm(client, new_batch, status=422, teacher_replacements={2: identifier})
    task_batch = await upload_task(client, deletion["task_id"], rows)
    assert (await client.get(P + f"/roster/teachers/deletions?search={task_batch['id']}")).json()["total"] == 1
    assert (await client.get(P + f"/roster/teachers/deletions?reason=change&search={deletion['task_id']}")).json()["total"] == 1
    assert (await client.get(P + "/roster/teachers/deletions?reason=depart")).json()["total"] == 0
    assert len(rows) == 1


@pytest.mark.asyncio
async def test_concurrent_task_batches_can_rebuild_each_original_only_once(client, test_db, test_engine, actors):
    if test_engine.dialect.name != "postgresql":
        pytest.skip("Independent PostgreSQL row locks required")
    import asyncio
    from sqlalchemy.ext.asyncio import async_sessionmaker
    from app.api.deps import AuthenticatedUser
    from app.core.errors import ApiError
    from app.db.models import School
    from app.schemas.roster import RosterConfirm
    from app.services.roster_workbench import confirm as confirm_service
    done = await confirm(client, await upload(client, [teacher("13900000001")]))
    identifier = done["report"]["imported"][0]["id"]
    deletion = await delete_teachers(client, [identifier])
    rows = await task_rows(client, deletion["task_id"])
    first = await upload_task(client, deletion["task_id"], [{**rows[0], "教师姓名": "并发修改甲"}])
    second = await upload_task(client, deletion["task_id"], [{**rows[0], "教师姓名": "并发修改乙"}])
    actor = AuthenticatedUser(actors["admin"].id, actors["school"].id, "qa-admin", ["SCHOOL_ADMIN"])
    factory = async_sessionmaker(test_engine, expire_on_commit=False)
    await test_db.commit()
    async def run(batch):
        async with factory() as db:
            try:
                await db.scalar(select(School.id).where(School.id == actor.school_id).with_for_update())
                await confirm_service(db, actor, UUID(batch["id"]), RosterConfirm(expected_revision=1, identity_confirmed=True, changes_confirmed=True))
                return 200
            except ApiError as error:
                await db.rollback()
                return error.status_code
    assert sorted(await asyncio.gather(run(first), run(second))) == [200, 409]
    assert await test_db.scalar(select(func.count()).select_from(TeacherProfile).where(TeacherProfile.id == UUID(identifier))) == 1
