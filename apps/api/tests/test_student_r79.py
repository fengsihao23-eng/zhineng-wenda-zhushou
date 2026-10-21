"""Real API/SQL acceptance for all six r79 student maintenance flows."""
import base64
import asyncio
import io
from datetime import datetime, timezone
from uuid import UUID, uuid4
import pytest
from openpyxl import load_workbook
from sqlalchemy import select, func
from app.api.deps import get_current_user
from app.core.database import get_db
from app.main import app
from app.core.security import get_password_hash, verify_password
from app.db.models import (Student, User, School, StudentYearClass, StudentExamPlacement, StudentWorkflowTask,
                           StudentExamScore, StudentSubjectScore, DiagnosisReport, ChatSession, ChatMessage, AgentRun, AuditLog)
from app.db.models.education import PrivateAsset, ReportAttachment
from app.services.roster_excel import STUDENT_COLUMNS, STUDENT_CHANGE_COLUMNS, STUDENT_LOCKED_FIELDS, STUDENT_PROMOTION_COLUMNS, workbook_bytes
from test_education_workbench import actors, add_scores, setup_exam, post, P  # noqa: F401
from test_roster_workflows import student, teacher, upload, confirm, classroom


async def task_create(client, kind, ids=(), values=None, headers=None, **extra):
    body = {"kind": kind, "student_ids": [str(s) for s in ids], **extra}
    if values is not None:
        columns = headers or (STUDENT_CHANGE_COLUMNS if kind == "archive" else STUDENT_PROMOTION_COLUMNS)
        body.update(filename="合成学生办理.xlsx", content_base64=base64.b64encode(workbook_bytes(columns, [[row.get(c, "") for c in columns] for row in values])).decode())
    return await post(client, "/roster/students/tasks", body)


async def task_prepare(client, task, **extra):
    defaults = {key: value for key, value in task.get("preparation", {}).items() if key not in {"prepared_by", "prepared_at"}}
    return await post(client, f"/roster/students/tasks/{task['id']}/prepare",
                      {**defaults, "expected_revision": task["revision"], "identity_confirmed_ids": [s["id"] for s in task["snapshots"]], **extra})


async def task_confirm(client, task, status=200):
    return await post(client, f"/roster/students/tasks/{task['id']}/confirm",
                      {"expected_revision": task["revision"], "confirmed": True, "identity_confirmed": True, "changes_confirmed": True}, status=status)


async def existing_fields(client, identifier):
    response = await client.get(P + f"/roster/students/records/{identifier}")
    assert response.status_code == 200, response.text
    return response.json()["fields"]


async def seed_student(client, account="new-student", name="合成新生", **values):
    batch = await upload(client, [student(account, name, **values)], "students")
    done = await confirm(client, batch, duplicate_decisions={r["row_number"]: True for r in batch["report"]["suspicious"]})
    return UUID(done["report"]["imported"][0]["id"])


@pytest.mark.asyncio
async def test_student_r79_export_has_31_exact_headers_and_eight_locked_columns(client, actors):
    sid = await seed_student(client)
    response = await client.get(P + "/roster/students/archive-template", params={"student_ids": str(sid)})
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    sheet = book.active
    assert [c.value for c in sheet[1]] == STUDENT_CHANGE_COLUMNS
    assert sheet.cell(2, 31).value == (await existing_fields(client, sid))["学生ID"]
    assert {STUDENT_CHANGE_COLUMNS[c.column - 1] for c in sheet[2] if c.protection.locked} == set(STUDENT_LOCKED_FIELDS)
    assert "mac地址 " not in STUDENT_COLUMNS and "SN码 " not in STUDENT_COLUMNS
    book.close()


@pytest.mark.asyncio
async def test_promotion_download_contains_only_selected_students_with_five_prefilled_columns(client, test_db, actors):
    first = await seed_student(client, "export-first", "名单甲", **{"年级（1-12）": "高中二年级", "班级号": "2班"})
    graduate = await seed_student(client, "export-graduate", "名单乙", **{"年级（1-12）": "高中三年级", "班级号": "3班"})
    await seed_student(client, "export-unselected", "未选中学生")
    await task_confirm(client, await task_prepare(client, await task_create(client, "status", [graduate], action="graduate")))
    fields = [await existing_fields(client, sid) for sid in [graduate, first]]
    response = await client.get(P + "/roster/students/promotion-template", params=[("student_ids", fields[0]["学生ID"]), ("student_ids", str(first))])
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    assert list(book.active.values) == [tuple(STUDENT_PROMOTION_COLUMNS)] + [
        (f["学生ID"], f["账号"], f["姓名"], f["年级（1-12）"], f["班级号"]) for f in fields]
    book.close()
    filtered = await client.get(P + "/roster/students/promotion-template", params={"all_filtered": True, "search": "export-graduate", "status": "graduated"})
    assert filtered.status_code == 200
    book = load_workbook(io.BytesIO(filtered.content))
    assert book.active.max_row == 2 and book.active.cell(2, 3).value == "名单乙"
    book.close()
    assert (await client.get(P + "/roster/students/promotion-template")).status_code == 422
    aliases = await client.get(P + "/roster/students/promotion-template", params=[("student_ids", str(first)), ("student_ids", fields[1]["学生ID"])])
    assert aliases.status_code == 422
    foreign = Student(id=uuid4(), school_id=actors["other_school"].id, name="跨校档案")
    test_db.add(foreign)
    await test_db.commit()
    assert (await client.get(P + "/roster/students/promotion-template", params={"student_ids": str(foreign.id)})).status_code == 422
    await task_confirm(client, await task_prepare(client, await task_create(client, "status", [first], action="depart")))
    assert (await client.get(P + "/roster/students/promotion-template", params={"student_ids": str(first)})).status_code == 422


@pytest.mark.asyncio
async def test_student_file_duplicates_cover_all_rows_and_explicit_keep_skip(client, actors):
    batch = await upload(client, [student("same-file-1"), student("same-file-2")], "students")
    assert {row["row_number"] for row in batch["report"]["suspicious"]} == {2, 3}
    await confirm(client, batch, status=422)
    done = await confirm(client, batch, duplicate_decisions={2: True, 3: False})
    assert done["report"]["success_count"] == 1 and done["report"]["skipped_rows"] == [3]
    # Blank grade/class are allowed but do not bypass a same-name reminder.
    batch = await upload(client, [student("blank-1", **{"年级（1-12）": "", "班级号": ""}), student("blank-2", **{"年级（1-12）": "", "班级号": ""})], "students")
    assert len(batch["report"]["suspicious"]) == 2


@pytest.mark.asyncio
async def test_unrestricted_student_optional_fields_and_long_utf8_initial_account_remain_usable(client, test_db, actors):
    account = "学生" * 50
    sid = await seed_student(client, account, "长账号验收", **{"年级（1-12）": "9" * 200, "班级号": "任意文本" * 100, "学号": "S" * 100, "手机号": "无格式校验" * 50})
    record = await test_db.get(Student, sid)
    assert len(record.profile_fields["班级号"]) == 400
    assert (await test_db.scalar(select(StudentYearClass).where(StudentYearClass.student_id == sid))).grade == "9" * 200
    response = await client.post("/api/v1/auth/login", json={"username": account, "password": account})
    assert response.status_code == 200 and response.json()["user"]["must_change_password"]
    wrong = await client.post("/api/v1/auth/login", json={"username": account, "password": account[:-1] + "错"})
    assert wrong.status_code == 401
    assert (await client.get(P + "/roster/students/records", params={"search": account})).json()["total"] == 1


@pytest.mark.asyncio
@pytest.mark.parametrize("account", ["student_basic", "password123"])
async def test_formal_student_account_format_is_unrestricted_in_production_but_first_reset_required(client, test_db, actors, monkeypatch, account):
    from app.core.config import settings
    await seed_student(client, account, "正式业务账号")
    monkeypatch.setattr(settings, "APP_ENV", "production")
    result = await client.post("/api/v1/auth/login", json={"username": account, "password": account})
    assert result.status_code == 200 and result.json()["user"]["must_change_password"]
    app.dependency_overrides.pop(get_current_user)
    blocked = await client.get("/api/v1/platform/student/dashboard", headers={"Authorization": "Bearer " + result.json()["access_token"]})
    assert blocked.status_code == 403 and "首次登录" in blocked.text


@pytest.mark.asyncio
async def test_archive_updates_ordinary_fields_and_clears_without_recreating_identity(client, test_db, actors):
    sid = await seed_student(client, 手机号="13800000001", 宿舍楼="A栋", 学号="OLD")
    record = await test_db.get(Student, sid)
    original_user = record.user_id
    fields = await existing_fields(client, sid)
    changed = {**fields, "手机号": "无需格式校验", "宿舍楼": "", "学号": "N" * 60}
    task = await task_create(client, "archive", values=[changed])
    assert {(r["field"], r["clear"]) for r in task["report"]["changes"]} == {("手机号", False), ("宿舍楼", True), ("学号", False)}
    await task_confirm(client, task, status=409)
    task = await task_prepare(client, (await client.get(P + f"/roster/students/tasks/{task['id']}")).json())
    done = await task_confirm(client, task)
    assert done["report"]["clear_count"] == 1 and done["report"]["success_count"] == 1
    assert record.id == sid and record.user_id == original_user and record.profile_fields["宿舍楼"] == ""
    assert record.student_no == "N" * 60
    assert (await task_confirm(client, task))["status"] == "succeeded"
    original = await client.get(P + f"/roster/students/tasks/{task['id']}/original")
    assert original.status_code == 200 and original.content


@pytest.mark.asyncio
async def test_archive_rejects_all_duplicate_ids_locked_fields_and_foreign_ids_atomically(client, test_db, actors):
    sid = await seed_student(client)
    fields = await existing_fields(client, sid)
    task = await task_create(client, "archive", values=[fields, {**fields, "姓名": "锁定姓名不可改"}])
    duplicate_rows = [r["row_number"] for r in task["report"]["issues"] if r["error_type"] == "ID重复"]
    assert duplicate_rows == [2, 3] and task["status"] == "invalid"
    await task_confirm(client, task, status=409)
    assert (await test_db.get(Student, sid)).name == fields["姓名"]
    foreign = Student(id=uuid4(), school_id=actors["other_school"].id, name="其他学校学生")
    test_db.add(foreign)
    await test_db.commit()
    bad = await task_create(client, "archive", values=[{**fields, "学生ID": str(foreign.id)}])
    assert bad["status"] == "invalid" and not bad["snapshots"]
    assert "其他学校学生" not in str(bad)


@pytest.mark.asyncio
async def test_identity_review_cannot_be_bypassed_and_approval_resumes_original_task(client, test_db, actors):
    sid = await seed_student(client)
    fields = await existing_fields(client, sid)
    task = await task_create(client, "archive", values=[{**fields, "宿舍楼": "B栋"}])
    task = await task_prepare(client, task, identity_confirmed_ids=[], identity_uncertain_ids=[str(sid)])
    assert task["status"] == "pending_review"
    task = await task_prepare(client, task, identity_confirmed_ids=[str(sid)], identity_uncertain_ids=[])
    assert task["status"] == "pending_review" and not task["report"]["ready_to_confirm"]
    await task_confirm(client, task, status=409)
    task = (await client.get(P + f"/roster/students/tasks/{task['id']}")).json()
    task = await post(client, f"/roster/students/tasks/{task['id']}/review", {"expected_revision": task["revision"], "student_ids": [str(sid)], "decision": "approved", "note": "已核对学籍号与原导入名单，确认为本人"})
    assert task["report"]["ready_to_confirm"]
    await task_confirm(client, task)
    assert (await test_db.get(Student, sid)).profile_fields["宿舍楼"] == "B栋"


@pytest.mark.asyncio
async def test_rejected_or_cancelled_task_and_concurrent_profile_change_never_write(client, test_db, actors):
    sid = await seed_student(client)
    fields = await existing_fields(client, sid)
    task = await task_create(client, "archive", values=[{**fields, "床号": "9"}])
    task = await task_prepare(client, task, identity_confirmed_ids=[], identity_uncertain_ids=[str(sid)])
    task = await post(client, f"/roster/students/tasks/{task['id']}/review", {"expected_revision": task["revision"], "student_ids": [str(sid)], "decision": "rejected", "note": "学籍号不符，无法确认本人"})
    await task_confirm(client, task, status=409)
    task = await task_create(client, "archive", values=[{**fields, "床号": "10"}])
    record = await test_db.get(Student, sid)
    record.profile_fields = {**record.profile_fields, "床号": "人工已更新"}
    await test_db.commit()
    task = await task_prepare(client, task)
    assert task["status"] == "invalid" and any(r["error_type"] == "原记录已变化" for r in task["report"]["issues"])
    cancelled = await post(client, f"/roster/students/tasks/{task['id']}/cancel", {"expected_revision": task["revision"]})
    assert cancelled["status"] == "cancelled" and record.profile_fields["床号"] == "人工已更新"


@pytest.mark.asyncio
async def test_name_duplicate_requires_confirmation_and_account_namespace_preserves_password_history(client, test_db, actors):
    sid = await seed_student(client, "old-student-account", "姓名甲")
    other = await seed_student(client, "other-student-account", "姓名乙")
    task = await task_create(client, "identity", [sid], field="姓名", new_value="姓名乙", reason="核对出生资料更正姓名")
    assert task["report"]["duplicates"][0]["student_id"] == str(other)
    task = await task_prepare(client, task)
    assert not task["report"]["ready_to_confirm"]
    task = await task_prepare(client, task, duplicate_name_confirmed=True)
    await task_confirm(client, task)
    done = await confirm(client, await upload(client, [teacher("13800000077")]))
    teacher_id = UUID(done["report"]["imported"][0]["id"])
    record = await test_db.get(Student, sid)
    user = await test_db.get(User, record.user_id)
    hashed, version = user.password_hash, user.password_version
    task = await task_create(client, "identity", [sid], field="账号", new_value="13800000077", reason="更正学校账号")
    assert not task["report"]["issues"]
    task = await task_prepare(client, task)
    await task_confirm(client, task)
    assert user.username == "13800000077" and user.password_hash == hashed and user.password_version == version + 1
    assert verify_password("old-student-account", user.password_hash)
    assert (await test_db.get(User, teacher_id)).username == "13800000077"
    history = (await client.get(P + "/roster/students/tasks", params={"search": "old-student-account"})).json()
    assert history["total"] == 2
    bad = await task_create(client, "identity", [sid], field="账号", new_value="other-student-account", reason="重复账号应拦截")
    assert bad["status"] == "invalid"


@pytest.mark.asyncio
async def test_graduation_and_restore_are_atomic_and_keep_original_account_id_and_exam_class(client, test_db, actors):
    sid = await seed_student(client, "graduate-student", "毕业学生", **{"年级（1-12）": "高中三年级", "班级号": "2班"})
    student_record = await test_db.get(Student, sid)
    original_user = student_record.user_id
    # Reuse real score fixtures, pointing their owner at this synthetic student.
    fixture_student = actors["student"]
    actors["student"] = type(fixture_student)(id=sid, user_id=original_user)
    exam = await setup_exam(client, actors)
    await add_scores(test_db, actors, exam)
    actors["student"] = fixture_student
    task = await task_create(client, "status", [sid], action="graduate")
    task = await task_prepare(client, task)
    await task_confirm(client, task)
    assert student_record.status == "graduated" and student_record.archived_at
    user = await test_db.get(User, original_user)
    assert user.status == "graduated"
    placement = await test_db.scalar(select(StudentExamPlacement).where(StudentExamPlacement.student_id == sid))
    assert placement.grade == "高中三年级"
    year = (await client.get(P + "/roster/students/context")).json()["current_year"]
    task = await task_create(client, "status", [sid], action="restore", placements={str(sid): {"grade": "高中二年级", "class_number": "3班"}})
    assert task["report"]["anomalies"]
    task = await task_prepare(client, task, anomaly_approvals={str(sid): "已核实同学段返校留级安排并审批通过"})
    await task_confirm(client, task)
    assert student_record.status == user.status == "active" and student_record.archived_at is None
    assert student_record.id == sid and student_record.user_id == original_user
    assert (await test_db.get(School, student_record.school_id)).current_academic_year == year
    assert placement.grade == "高中三年级"
    assert await test_db.scalar(select(func.count()).select_from(StudentExamScore).where(StudentExamScore.student_id == sid)) == 1
    await task_confirm(client, task)


@pytest.mark.asyncio
async def test_mixed_status_list_and_cross_segment_restore_are_rejected_without_partial_effects(client, test_db, actors):
    left = await seed_student(client, "departed-student", "离校学生")
    active = await seed_student(client, "active-student", "在读学生")
    task = await task_prepare(client, await task_create(client, "status", [left], action="depart"))
    await task_confirm(client, task)
    bad = await task_create(client, "status", [left, active], action="depart")
    assert bad["status"] == "invalid"
    await task_confirm(client, bad, status=409)
    assert (await test_db.get(Student, active)).status == "active"
    bad = await task_create(client, "status", [left], action="restore", placements={str(left): {"grade": "初中一年级", "class_number": "1班"}})
    assert bad["status"] == "invalid" and "同学段" in str(bad["report"]["issues"])


@pytest.mark.asyncio
async def test_next_year_requires_verified_graduation_and_complete_active_roster_before_switch(client, test_db, actors):
    fixture = await test_db.get(Student, actors["student"].id)
    fixture.profile_fields = {"年级（1-12）": "高中一年级", "班级号": "1班"}
    await test_db.commit()
    regular = await seed_student(client, "promote-regular", "升级学生", **{"年级（1-12）": "高中二年级"})
    graduate = await seed_student(client, "promote-graduate", "应届学生", **{"年级（1-12）": "高中三年级"})
    fields = [await existing_fields(client, sid) for sid in [fixture.id, regular]]
    values = [{"学生ID": f["学生ID"], "姓名": f["姓名"], "账号": f["账号"], "目标年级": "高中二年级" if f["学生ID"] == str(fixture.platform_id) else "高中三年级", "目标班级号": "2班"} for f in fields]
    task = await task_create(client, "promotion", values=values, mode="next")
    assert len(task["report"]["missing"]) == 1
    task = await task_prepare(client, task, graduation_confirmed=True, graduation_student_ids=[str(graduate)])
    assert task["report"]["pending_graduations"] and not task["report"]["ready_to_confirm"]
    school = await test_db.get(School, fixture.school_id)
    original_year = school.current_academic_year
    graduation = await task_prepare(client, await task_create(client, "status", [graduate], action="graduate"))
    await task_confirm(client, graduation)
    refreshed = (await client.get(P + f"/roster/students/tasks/{task['id']}")).json()
    assert refreshed["revision"] == task["revision"]
    assert refreshed["report"]["graduation_completed"][0]["student_id"] == str(graduate)
    assert not refreshed["report"]["graduation_reviewed"]
    task = await task_prepare(client, task)
    assert not task["report"]["ready_to_confirm"]
    task = await task_prepare(client, task, graduation_review_digest=task["report"]["graduation_review_digest"])
    assert task["report"]["ready_to_confirm"] and not task["report"]["missing"]
    done = await task_confirm(client, task)
    assert done["status"] == "succeeded" and school.current_academic_year != original_year
    assert (await test_db.get(Student, graduate)).status == "graduated"
    assert await test_db.scalar(select(func.count()).select_from(StudentYearClass).where(StudentYearClass.student_id == regular)) == 2


@pytest.mark.asyncio
@pytest.mark.parametrize("mode", ["current", "next"])
async def test_graduated_student_can_repeat_via_promotion_after_graduation_review_without_changing_identity_or_exam_class(client, test_db, actors, mode):
    fixture = await test_db.get(Student, actors["student"].id)
    fixture.profile_fields = {"年级（1-12）": "高中一年级", "班级号": "1班"}
    await test_db.commit()
    sid = await seed_student(client, "promotion-repeater", "毕业后复读学生", **{"年级（1-12）": "高中三年级", "班级号": "2班"})
    record = await test_db.get(Student, sid)
    user = await test_db.get(User, record.user_id)
    user_id, password_hash = record.user_id, user.password_hash
    original_actor = actors["student"]
    actors["student"] = type(original_actor)(id=sid, user_id=user_id)
    exam = await setup_exam(client, actors)
    await add_scores(test_db, actors, exam)
    actors["student"] = original_actor
    await task_confirm(client, await task_prepare(client, await task_create(client, "status", [sid], action="graduate")))
    school = await test_db.get(School, fixture.school_id)
    current_year = school.current_academic_year
    fields = await existing_fields(client, sid)
    values = [{"学生ID": fields["学生ID"], "姓名": fields["姓名"], "账号": fields["账号"], "目标年级": "复读班", "目标班级号": "3班"}]
    if mode == "next":
        regular = await existing_fields(client, fixture.id)
        values.append({"学生ID": regular["学生ID"], "姓名": regular["姓名"], "账号": regular["账号"], "目标年级": "高中二年级", "目标班级号": "1班"})
    task = await task_create(client, "promotion", values=values, mode=mode)
    assert not task["report"]["issues"] and not task["report"]["missing"]
    completed = task["report"]["graduation_completed"]
    assert len(completed) == 1
    assert (completed[0]["platform_id"], completed[0]["graduation_grade"], completed[0]["graduation_class_number"], completed[0]["repeat_grade"], completed[0]["repeat_class_number"]) == (fields["学生ID"], "高中三年级", "2班", "复读班", "3班")
    task = await task_prepare(client, task, identity_confirmed_ids=[], graduation_confirmed=True)
    assert task["report"]["graduation_reviewed"] and not task["report"]["identity_confirmed"]
    assert record.status == user.status == "graduated" and record.archived_at
    task = await task_prepare(client, task, anomaly_approvals={str(sid): "已核实同学段毕业后复读安排，审批通过"})
    assert task["report"]["ready_to_confirm"]
    done = await task_confirm(client, task)
    assert done["status"] == "succeeded"
    assert record.id == sid and record.user_id == user_id and user.password_hash == password_hash
    assert record.status == user.status == "active" and record.archived_at is None
    assert (await existing_fields(client, sid))["学生ID"] == fields["学生ID"]
    assert record.profile_fields["年级（1-12）"] == "复读班" and record.profile_fields["班级号"] == "3班"
    placement = await test_db.scalar(select(StudentExamPlacement).where(StudentExamPlacement.student_id == sid))
    assert placement.grade == "高中三年级" and placement.class_number == "2班"
    assert await test_db.scalar(select(func.count()).select_from(StudentExamScore).where(StudentExamScore.student_id == sid)) == 1
    if mode == "next":
        year_record = await test_db.scalar(select(StudentYearClass).where(StudentYearClass.student_id == sid, StudentYearClass.academic_year == current_year))
        assert year_record.status == "graduated" and year_record.grade == "高中三年级"
    else:
        repeat = {**values[0], "目标班级号": "4班"}
        followup = await task_create(client, "promotion", values=[repeat], mode="current")
        assert followup["report"]["graduation_completed"][0]["graduation_grade"] == "高中三年级"
        assert followup["report"]["graduation_completed"][0]["graduation_class_number"] == "2班"


@pytest.mark.asyncio
async def test_promotion_does_not_restore_a_graduate_into_a_different_school_segment(client, test_db, actors):
    sid = await seed_student(client, "promotion-cross-segment", "跨学段复读拦截", **{"年级（1-12）": "初中三年级"})
    await task_confirm(client, await task_prepare(client, await task_create(client, "status", [sid], action="graduate")))
    fields = await existing_fields(client, sid)
    task = await task_create(client, "promotion", values=[{"学生ID": fields["学生ID"], "姓名": fields["姓名"], "账号": fields["账号"], "目标年级": "高中一年级", "目标班级号": "1班"}])
    assert task["status"] == "invalid" and "同学段" in str(task["report"]["issues"])
    await task_confirm(client, task, status=409)
    assert (await test_db.get(Student, sid)).status == "graduated"


@pytest.mark.asyncio
async def test_next_year_retained_graduating_grade_student_needs_no_graduation_exclusion_selection(client, test_db, actors):
    fixture = await test_db.get(Student, actors["student"].id)
    fixture.profile_fields = {"年级（1-12）": "高中三年级", "班级号": "1班"}
    await test_db.commit()
    fields = await existing_fields(client, fixture.id)
    task = await task_create(client, "promotion", values=[{"学生ID": fields["学生ID"], "姓名": fields["姓名"], "账号": fields["账号"], "目标年级": "高中三年级", "目标班级号": "2班"}], mode="next")
    assert not task["report"]["graduation_completed"] and not task["report"]["pending_graduations"]
    assert not task["report"]["graduation_exclusions_missing"]
    task = await task_prepare(client, task, graduation_confirmed=True, anomaly_approvals={str(fixture.id): "已核实本学年不毕业，留级安排审批通过"})
    assert task["report"]["ready_to_confirm"]
    assert (await task_confirm(client, task))["status"] == "succeeded"


@pytest.mark.asyncio
async def test_deletion_cleans_personal_descendants_recalculates_only_participating_exams_and_keeps_snapshots(client, test_db, actors):
    sid = actors["student"].id
    await classroom(client)
    record = await test_db.get(Student, sid)
    record.profile_fields = {"年级（1-12）": "高中一年级", "班级号": "1班"}
    await test_db.commit()
    survivor = await seed_student(client, "survivor-student", "保留学生")
    other_class = await seed_student(client, "other-class-student", "其他班级学生", 班级号="2班")
    other_grade = await seed_student(client, "other-grade-student", "其他年级学生", **{"年级（1-12）": "高中二年级"})
    exam = await setup_exam(client, actors)
    await add_scores(test_db, actors, exam)
    survivor_score = StudentExamScore(school_id=record.school_id, student_id=survivor, exam_id=UUID(exam["id"]), total_score=450, class_rank=8, grade_rank=9)
    survivor_subject = StudentSubjectScore(school_id=record.school_id, student_id=survivor, exam_id=UUID(exam["id"]), subject_id=actors["subject"].id, score=70, full_score=100, class_rank=8, grade_rank=9)
    same_grade_score = StudentExamScore(school_id=record.school_id, student_id=other_class, exam_id=UUID(exam["id"]), total_score=530, class_rank=73, grade_rank=74)
    other_grade_score = StudentExamScore(school_id=record.school_id, student_id=other_grade, exam_id=UUID(exam["id"]), total_score=600, class_rank=83, grade_rank=84)
    unrelated = await post(client, "/school/exams", {"external_exam_id": "unrelated-exam", "name": "无关考试", "exam_type": "月考", "start_date": "2026-09-10"})
    untouched = StudentExamScore(school_id=record.school_id, student_id=survivor, exam_id=UUID(unrelated["id"]), total_score=999, class_rank=91, grade_rank=92)
    session = ChatSession(id=uuid4(), school_id=record.school_id, student_id=sid)
    test_db.add_all([survivor_score, survivor_subject, same_grade_score, other_grade_score, untouched, session])
    await test_db.flush()
    message = ChatMessage(session_id=session.id, role="user", content="仅属于待删除学生的问答")
    run = AgentRun(session_id=session.id, school_id=record.school_id, student_id=sid, query="合成个人分析", entitlement_level="BASIC", status="succeeded")
    report = DiagnosisReport(id=uuid4(), school_id=record.school_id, student_id=sid, exam_id=UUID(exam["id"]), report_type="diagnosis", structured_json={}, generated_at=datetime.now(timezone.utc))
    survivor_report = DiagnosisReport(school_id=record.school_id, student_id=survivor, exam_id=UUID(exam["id"]), report_type="diagnosis", structured_json={"original": "保留原报告"}, generated_at=datetime.now(timezone.utc))
    asset = PrivateAsset(id=uuid4(), school_id=record.school_id, filename="合成个人报告.txt", media_type="text/plain", byte_size=4, content_hash="synthetic", storage_key=str(uuid4()), content=b"data", created_by=actors["admin"].id)
    test_db.add_all([message, run, report, survivor_report, asset])
    await test_db.flush()
    test_db.add(ReportAttachment(school_id=record.school_id, report_id=report.id, asset_id=asset.id))
    await test_db.commit()
    original_user = record.user_id
    task = await task_create(client, "deletion", [sid], reason="已对照原导入名单核实为误导入")
    assert any(row["type"] == "student_exam_scores" for row in task["report"]["affected"])
    assert [row["id"] for row in task["report"]["affected_exams"]] == [exam["id"]]
    task = await task_prepare(client, task, records_verified_ids=[str(sid)])
    done = await task_confirm(client, task)
    assert done["report"]["disposition"]["deleted_students"] == 1
    assert await test_db.get(Student, sid) is None and await test_db.get(User, original_user) is None
    assert await test_db.scalar(select(ChatSession.id).where(ChatSession.id == session.id)) is None
    assert await test_db.scalar(select(PrivateAsset.id).where(PrivateAsset.id == asset.id)) is None
    await test_db.refresh(survivor_score)
    await test_db.refresh(untouched)
    await test_db.refresh(survivor_report)
    assert survivor_score.class_rank == 1 and survivor_score.grade_rank == 2
    await test_db.refresh(same_grade_score)
    await test_db.refresh(other_grade_score)
    assert same_grade_score.class_rank == 73 and same_grade_score.grade_rank == 1
    assert other_grade_score.class_rank == 83 and other_grade_score.grade_rank == 84
    assert untouched.class_rank == 91 and untouched.grade_rank == 92
    assert survivor_report.structured_json["original"] == "保留原报告"
    assert survivor_report.structured_json["recalculated_statistics"]["class_rank"] == 1
    retained = await test_db.get(StudentWorkflowTask, UUID(task["id"]))
    assert retained.snapshots[0]["id"] == str(sid) and retained.snapshots[0]["fields"]["账号"]


@pytest.mark.asyncio
async def test_new_workflows_are_tenant_scoped_school_admin_only_and_old_shortcuts_blocked(client, actors):
    sid = await seed_student(client)
    task = await task_create(client, "identity", [sid], field="姓名", new_value="姓名更正", reason="已核对资料")
    for role in ["teacher", "student", "qa"]:
        actors["switch"](role)
        await post(client, f"/roster/students/tasks/{task['id']}/prepare", {"expected_revision": task["revision"]}, status=403)
    actors["switch"]("foreign")
    assert (await client.get(P + f"/roster/students/tasks/{task['id']}")).status_code == 404
    actors["switch"]()
    result = await post(client, f"/roster/students/records/{sid}/actions", {"action": "delete", "confirmed": True}, status=422)
    assert result["error"]["code"] == "STUDENT_WORKFLOW_REQUIRED"


@pytest.mark.asyncio
async def test_concurrent_student_confirmation_commits_one_batch_once(client, test_db, test_engine, actors):
    if test_engine.dialect.name != "postgresql":
        pytest.skip("Independent PostgreSQL school and task locks required")
    from sqlalchemy.ext.asyncio import async_sessionmaker
    sid = await seed_student(client, "concurrent-student", "并发验收学生")
    fields = await existing_fields(client, sid)
    task = await task_prepare(client, await task_create(client, "archive", values=[{**fields, "床号": "88"}]))
    factory = async_sessionmaker(test_engine, expire_on_commit=False)
    original = app.dependency_overrides[get_db]

    async def independent_db():
        async with factory() as session:
            yield session

    app.dependency_overrides[get_db] = independent_db
    try:
        body = {"expected_revision": task["revision"], "confirmed": True, "identity_confirmed": True, "changes_confirmed": True}
        responses = await asyncio.gather(*(client.post(P + f"/roster/students/tasks/{task['id']}/confirm", json=body) for _ in range(2)))
        assert [response.status_code for response in responses] == [200, 200]
        assert all(response.json()["report"]["success_count"] == 1 for response in responses)
    finally:
        app.dependency_overrides[get_db] = original
    await test_db.refresh(await test_db.get(Student, sid))
    assert (await test_db.get(Student, sid)).profile_fields["床号"] == "88"
    assert await test_db.scalar(select(func.count()).select_from(AuditLog).where(
        AuditLog.action == "education.student.archive.confirm", AuditLog.resource_id == UUID(task["id"]))) == 1


@pytest.mark.asyncio
async def test_erroneous_deletion_includes_shared_web_mini_account_data_and_recovers_old_blocked_task(client, test_db, actors):
    from app.db.models.answer_workbench import (AnswerSession, AnswerTicket, AnswerOperation,
                                               AnswerTrial, AnswerDailyUsage, AnswerPresence)
    from app.services.answer_capacity import day_at
    models = (AnswerSession, AnswerTicket, AnswerOperation, AnswerTrial, AnswerDailyUsage, AnswerPresence)
    selected = [await seed_student(client, "delete-mini-one", "误导入甲"),
                await seed_student(client, "delete-mini-two", "误导入乙")]
    survivor = await seed_student(client, "keep-mini-account", "保留问答学生")
    people = [await test_db.get(Student, sid) for sid in [*selected, survivor]]
    foreign = User(id=uuid4(), school_id=actors["other_school"].id, username="foreign-mini-user", password_hash="unused")
    test_db.add(foreign)
    await test_db.flush()
    owners = [(person.school_id, person.user_id) for person in people] + [(foreign.school_id, foreign.id)]
    for school_id, user_id in owners:
        common = {"school_id": school_id, "user_id": user_id}
        test_db.add_all([
            AnswerSession(**common, body={"messages": []}),
            AnswerTicket(**common, body={"question": "合成留言", "events": []}),
            AnswerOperation(**common, action="合成个人操作", body={}),
            AnswerTrial(**common, scope="prompt:高中", body={}),
            AnswerDailyUsage(**common, day=day_at(), used=1),
            AnswerPresence(**common, client_id=uuid4(), expires_at=datetime.now(timezone.utc)),
        ])
    await test_db.commit()
    task = await task_create(client, "deletion", selected, reason="核对原导入文件，确认两条均为误导入")
    assert not task["report"]["issues"] and not task["report"]["deletion_barriers"]
    assert not task["report"]["ready_to_confirm"]
    assert {model.__tablename__ for model in models} <= {row["type"] for row in task["report"]["affected"]}
    # A previously saved failure must refresh to the current rules without
    # silently substituting the user's required identity/record confirmations.
    saved = await test_db.get(StudentWorkflowTask, UUID(task["id"]))
    saved.status = "invalid"
    saved.report = {**saved.report, "issues": [{"message": "旧版误判：问答记录属于共享业务"}],
                    "deletion_barriers": ["answer_sessions", "answer_daily_usage", "answer_presence"]}
    await test_db.commit()
    response = await client.get(P + f"/roster/students/tasks/{task['id']}")
    assert response.status_code == 200
    task = response.json()
    assert task["status"] == "ready" and not task["report"]["issues"]
    assert not task["report"]["ready_to_confirm"]
    task = await task_prepare(client, task, records_verified_ids=[str(sid) for sid in selected])
    assert task["report"]["ready_to_confirm"]
    done = await task_confirm(client, task)
    assert done["report"]["disposition"]["deleted_students"] == 2
    replay = await task_confirm(client, task)
    assert replay["id"] == done["id"] and replay["status"] == "succeeded"
    assert replay["revision"] == done["revision"] and replay["report"] == done["report"]
    for model in models:
        assert done["report"]["disposition"]["personal_data_counts"][model.__tablename__] == 2
        for school_id, user_id in owners[:2]:
            assert await test_db.scalar(select(func.count()).select_from(model).where(
                model.school_id == school_id, model.user_id == user_id)) == 0
        for school_id, user_id in owners[2:]:
            assert await test_db.scalar(select(func.count()).select_from(model).where(
                model.school_id == school_id, model.user_id == user_id)) == 1
    assert await test_db.get(Student, survivor) is not None
    assert len(done["snapshots"]) == 2


@pytest.mark.asyncio
async def test_erroneous_deletion_still_blocks_account_ownership_of_unrelated_shared_files(client, test_db, actors):
    sid = await seed_student(client, "shared-file-owner", "共享文件归属学生")
    person = await test_db.get(Student, sid)
    asset = PrivateAsset(school_id=person.school_id, filename="共享资料.pdf", media_type="application/pdf",
                         byte_size=4, content_hash="shared", storage_key="shared-owner-test", content=b"test",
                         created_by=person.user_id)
    test_db.add(asset)
    await test_db.commit()
    task = await task_create(client, "deletion", [sid], reason="合成误导入核验")
    assert "education_assets" in task["report"]["deletion_barriers"]
    task = await task_prepare(client, task, records_verified_ids=[str(sid)])
    assert not task["report"]["ready_to_confirm"]
    await task_confirm(client, task, status=409)
    assert await test_db.get(Student, sid) is not None
    assert await test_db.get(PrivateAsset, asset.id) is not None
