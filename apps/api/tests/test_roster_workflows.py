"""Executable acceptance for all four flows in the supplied 0922 HTML."""
import base64
import io
from uuid import UUID, uuid4
import pytest
from openpyxl import load_workbook
from sqlalchemy import select, func
from app.main import app
from app.api.deps import get_current_user
from app.db.models import User, Student, StudentExamScore, AuditLog, TeachingAssignment
from app.db.models.roster import TeacherProfile
from app.services.roster_excel import TEACHER_COLUMNS, STUDENT_COLUMNS, workbook_bytes
from test_education_workbench import actors, post, setup_exam, add_scores, P  # noqa: F401


def teacher(account="teacher-a", name="合成教师甲", **values):
    return {"教师账号": account, "教师姓名": name, "状态": "正常", **values}


def student(account="student-a", name="合成学生甲", **values):
    return {"账号": account, "姓名": name, "年级（1-12）": "高中一年级", "班级号": "1班", **values}


async def upload(client, rows, kind="teachers", columns=None):
    columns = columns or (TEACHER_COLUMNS if kind == "teachers" else STUDENT_COLUMNS)
    content = workbook_bytes(columns, [[row.get(c, "") for c in columns] for row in rows])
    return await post(client, f"/roster/{kind}/imports", {"filename": "合成资料.xlsx", "content_base64": base64.b64encode(content).decode()})


async def confirm(client, batch, status=200, **values):
    return await post(client, f"/roster/imports/{batch['id']}/confirm", {"expected_revision": batch["revision"], **values}, status=status)


async def lifecycle(client, kind, identifier, action="delete", status=200):
    return await post(client, f"/roster/{kind}/records/{identifier}/actions", {"action": action, "confirmed": True}, status=status)


async def classroom(client):
    return await post(client, "/school/classes", {"external_class_id": "10.1", "name": "高中一年级1班"})


@pytest.mark.asyncio
async def test_template_contract_and_file_level_rejection(client, test_db, actors):
    for kind, columns in [("teachers", TEACHER_COLUMNS), ("students", STUDENT_COLUMNS)]:
        response = await client.get(P + f"/roster/{kind}/template")
        book = load_workbook(io.BytesIO(response.content))
        assert [c.value for c in book.active[1]] == columns
        book.close()
    student_errors = await client.get(P + "/roster/students/error-template")
    error_book = load_workbook(io.BytesIO(student_errors.content))
    assert [c.value for c in error_book.active[1]] == ["行号", "错误类型", "错误字段", "姓名", "账号", "错误说明", "原始内容", "处理建议"]
    assert error_book.active.max_row == 8 and error_book.active.max_column == 8
    error_book.close()
    before = await test_db.scalar(select(func.count()).select_from(User))
    columns = TEACHER_COLUMNS.copy()
    columns[0], columns[1] = columns[1], columns[0]
    batch = await upload(client, [teacher("", "")], columns=columns)
    assert batch["status"] == "invalid" and batch["rows"] == []
    assert len(batch["report"]["issues"]) == 1
    assert batch["report"]["issues"][0]["row_number"] is None
    await confirm(client, batch, status=409)
    assert await test_db.scalar(select(func.count()).select_from(User)) == before


@pytest.mark.asyncio
async def test_light_preflight_merges_row_errors_and_keeps_masked_phone(client, test_db, actors):
    batch = await upload(client, [teacher("", "", 状态="离职"), teacher("same"), teacher("same"), teacher("other", 手机号码="131****7000", 任课年级班级="体育:99.999;")])
    issues = batch["report"]["issues"]
    assert len(issues) == 2
    assert issues[0]["fields"] == "教师姓名、教师账号、状态"
    assert "第 3 行" in issues[1]["message"]
    download = await client.get(P + f"/roster/imports/{batch['id']}/errors")
    book = load_workbook(io.BytesIO(download.content))
    assert book.active.max_row == 3 and book.active.max_column == 8
    book.close()
    await confirm(client, batch, status=409)
    assert await test_db.scalar(select(User).where(User.username == "other")) is None
    good = await upload(client, [teacher("other", 手机号码="131****7000", 任课年级班级="体育:99.999;")])
    assert good["report"]["ok"]
    done = await confirm(client, good)
    assert done["report"]["success_count"] == 1
    assert done["report"]["unresolved_teaching"]
    profile = await test_db.get(TeacherProfile, UUID(done["report"]["imported"][0]["id"]))
    assert profile.fields["手机号码"] == "131****7000" and profile.fields["学工号"] == ""
    duplicate = await upload(client, [teacher("other")])
    assert duplicate["status"] == "invalid"


@pytest.mark.asyncio
async def test_teacher_uniqueness_only_within_teachers_and_login_type(client, test_db, actors):
    await classroom(client)
    s = await confirm(client, await upload(client, [student("shared-account")], "students"))
    t = await confirm(client, await upload(client, [teacher("shared-account")]))
    assert s["report"]["success_count"] == t["report"]["success_count"] == 1
    ambiguous = await client.post("/api/v1/auth/login", json={"username": "shared-account", "password": "shared-account"})
    assert ambiguous.status_code == 409
    assert {identity["account_type"] for identity in ambiguous.json()["error"]["details"]["identities"]} == {"teacher", "student"}
    for kind, role in [("teacher", "TEACHER"), ("student", "STUDENT")]:
        response = await client.post("/api/v1/auth/login", json={"username": "shared-account", "password": "shared-account", "account_type": kind})
        assert response.status_code == 200, response.text
        assert role in response.json()["user"]["roles"]
        assert response.json()["user"]["must_change_password"]


@pytest.mark.asyncio
async def test_suspects_require_each_choice_and_never_silently_disappear(client, test_db, actors):
    await confirm(client, await upload(client, [teacher("old", 任课年级班级="数学:10.1;")]))
    batch = await upload(client, [teacher("keep", 任课年级班级="数学:10.2;", 手机号码="different"), teacher("skip", 任课年级班级="数学:10.3;"), teacher("blank"), teacher("other-grade", 任课年级班级="数学:11.1;")])
    assert [s["row_number"] for s in batch["report"]["suspicious"]] == [2, 3]
    assert all(s["keep"] for s in batch["report"]["suspicious"])
    await confirm(client, batch, status=422)
    assert await test_db.scalar(select(User).where(User.username == "keep")) is None
    result = await confirm(client, batch, duplicate_decisions={2: True, 3: False})
    assert result["report"]["success_count"] == 3 and result["report"]["skipped_rows"] == [3]
    replay = await confirm(client, batch, duplicate_decisions={2: True, 3: False})
    assert replay["report"] == result["report"]
    await confirm(client, batch, duplicate_decisions={2: False, 3: True}, status=409)
    assert await test_db.scalar(select(User).where(User.username == "skip")) is None


@pytest.mark.asyncio
async def test_confirm_revalidates_before_any_write(client, test_db, actors):
    first = await upload(client, [teacher("raced"), teacher("unwritten", "另一教师")])
    await confirm(client, await upload(client, [teacher("raced")]))
    await confirm(client, first, status=409)
    assert await test_db.scalar(select(User).where(User.username == "unwritten")) is None


@pytest.mark.asyncio
async def test_roles_grants_and_delete_preserves_history_reimport(client, test_db, actors):
    await classroom(client)
    body = teacher("scoped-teacher", 说明="班主任(高中一年级1班)\n任课教师", 任课年级班级="数学:10.1;", 班主任年级班级="10.1")
    result = await confirm(client, await upload(client, [body]))
    identifier = result["report"]["imported"][0]["id"]
    profile = await test_db.get(TeacherProfile, UUID(identifier))
    assert "HOMEROOM_TEACHER" in profile.duties
    grant = await test_db.scalar(select(TeachingAssignment).where(TeachingAssignment.teacher_user_id == UUID(identifier)))
    assert grant is not None
    deletion = await lifecycle(client, "teachers", identifier)
    assert (await test_db.get(User, UUID(identifier))).status == "deleted"
    assert grant.teacher_user_id == UUID(identifier)
    snapshot = (await client.get(P + f"/roster/teachers/deletions/{deletion['deletion_id']}")).json()
    assert snapshot["fields"] == profile.fields
    assert snapshot["teaching_history"][0]["id"] == str(grant.id)
    assert (await client.get(P + "/roster/teachers/records?search=scoped-teacher")).json()["total"] == 0
    changed = await confirm(client, await upload(client, [teacher("scoped-teacher", "变更后的教师")]))
    assert changed["report"]["success_count"] == changed["report"]["changed_count"] == 1
    snapshot_after = (await client.get(P + f"/roster/teachers/deletions/{deletion['deletion_id']}")).json()
    assert snapshot_after["fields"] == profile.fields
    assert snapshot_after["reimported_user_id"] == changed["report"]["imported"][0]["id"]
    assert await test_db.scalar(select(func.count()).select_from(AuditLog).where(AuditLog.action == "education.roster.teachers.delete")) == 1


@pytest.mark.asyncio
async def test_first_login_forced_change_and_password_revokes_old_tokens(client, actors):
    result = await confirm(client, await upload(client, [teacher("first-login")]))
    app.dependency_overrides.pop(get_current_user)
    response = await client.post("/api/v1/auth/login", json={"username": "first-login", "password": "first-login"})
    tokens = response.json()
    headers = {"Authorization": "Bearer " + tokens["access_token"]}
    assert (await client.get("/api/v1/auth/me", headers=headers)).json()["must_change_password"]
    blocked = await client.get(P + "/school/classes", headers=headers)
    assert blocked.status_code == 403 and blocked.headers["X-Error-Code"] == "PASSWORD_CHANGE_REQUIRED"
    bad = await client.post("/api/v1/auth/change-password", headers=headers, json={"current_password": "first-login", "new_password": "first-login"})
    assert bad.status_code == 422
    changed = await client.post("/api/v1/auth/change-password", headers=headers, json={"current_password": "first-login", "new_password": "changed-secret-0922"})
    assert changed.status_code == 200, changed.text
    new_headers = {"Authorization": "Bearer " + changed.json()["access_token"]}
    assert (await client.get(P + "/school/classes", headers=new_headers)).status_code == 200
    assert (await client.get("/api/v1/auth/me", headers=headers)).status_code == 401
    assert (await client.post("/api/v1/auth/refresh", json={"refresh_token": tokens["refresh_token"]})).status_code == 401
    assert result["report"]["success_count"] == 1


@pytest.mark.asyncio
async def test_principal_no_class_import_gets_school_read_only(client, actors):
    await classroom(client)
    result = await confirm(client, await upload(client, [teacher("principal-login", 是否校长="是")]))
    assert result["report"]["success_count"] == 1
    app.dependency_overrides.pop(get_current_user)
    login = (await client.post("/api/v1/auth/login", json={"username": "principal-login", "password": "principal-login"})).json()
    headers = {"Authorization": "Bearer " + login["access_token"]}
    changed = await client.post("/api/v1/auth/change-password", headers=headers, json={"current_password": "principal-login", "new_password": "principal-updated-0922"})
    headers = {"Authorization": "Bearer " + changed.json()["access_token"]}
    assert (await client.get(P + "/school/classes", headers=headers)).json()["total"] == 1
    assert (await client.get(P + "/roster/teachers/records", headers=headers)).status_code == 200
    denied = await client.post(P + "/school/classes", headers=headers, json={"name": "禁止修改", "external_class_id": "forbidden"})
    assert denied.status_code == 403


@pytest.mark.asyncio
async def test_students_keep_profile_class_without_dictionary_match(client, test_db, actors):
    missing_class = await upload(client, [student("missing-class")], "students")
    assert missing_class["status"] == "ready" and missing_class["report"]["ok"]
    imported_without_dictionary = await confirm(client, missing_class)
    first = await test_db.get(Student, UUID(imported_without_dictionary["report"]["imported"][0]["id"]))
    assert first.user_id and first.student_no is None and first.class_id is None
    assert first.profile_fields["年级（1-12）"] == "高中一年级"
    await classroom(client)
    done = await confirm(client, await upload(client, [student("new-student", "新学生甲", 手机号="no-format-check"), student("second-student", "新学生乙", 学籍号="same-is-not-identity")], "students"))
    assert done["report"]["success_count"] == 2
    for row in done["report"]["imported"]:
        item = await test_db.get(Student, UUID(row["id"]))
        assert item.user_id and item.student_no is None and item.class_id
    duplicate = await upload(client, [student("new-student")], "students")
    assert duplicate["status"] == "invalid"


@pytest.mark.asyncio
async def test_student_suspected_duplicate_requires_row_confirmation(client, test_db, actors):
    original = await confirm(client, await upload(client, [student("original-account", "同名学生")], "students"))
    original_id = UUID(original["report"]["imported"][0]["id"])
    batch = await upload(client, [student("replacement-account", "同名学生")], "students")
    assert batch["status"] == "ready"
    assert batch["report"]["suspicious"][0]["row_number"] == 2
    assert batch["report"]["suspicious"][0]["matches"][0]["id"] == str(original_id)
    await confirm(client, batch, status=422)
    done = await confirm(client, batch, duplicate_decisions={2: True})
    assert done["report"]["success_count"] == 1 and done["report"]["skipped_count"] == 0


@pytest.mark.asyncio
async def test_student_delete_with_scores_blocks_but_graduation_keeps_history(client, test_db, actors):
    exam = await setup_exam(client, actors)
    await add_scores(test_db, actors, exam)
    identifier = str(actors["student"].id)
    await lifecycle(client, "students", identifier, status=409)
    result = await lifecycle(client, "students", identifier, action="graduate")
    assert result["status"] == "graduated"
    item = await test_db.get(Student, actors["student"].id)
    assert item.status == "graduated"
    assert (await test_db.get(User, item.user_id)).status == "graduated"
    assert await test_db.scalar(select(func.count()).select_from(StudentExamScore).where(StudentExamScore.student_id == item.id)) == 1
    assert (await lifecycle(client, "students", identifier, action="graduate"))["status"] == "graduated"


@pytest.mark.asyncio
async def test_student_change_delete_reimport(client, test_db, actors):
    await classroom(client)
    batch = await upload(client, [student("moved-student"), student("kept-student")], "students")
    done = await confirm(client, batch)
    identifier = done["report"]["imported"][0]["id"]
    assert "parent_binding_count" not in done["report"]
    await lifecycle(client, "students", identifier)
    assert await test_db.get(Student, UUID(identifier)) is None
    await post(client, "/school/classes", {"external_class_id": "10.2", "name": "高中一年级2班"})
    changed = await confirm(client, await upload(client, [student("moved-student", 班级号="2班")], "students"))
    item = await test_db.get(Student, UUID(changed["report"]["imported"][0]["id"]))
    assert item.external_class_id == "10.2"
    assert (await client.get(P + "/roster/students/" + identifier + "/parents")).status_code == 404
    await post(client, "/roster/parents/" + identifier + "/unbind", {}, status=404)


@pytest.mark.asyncio
async def test_departure_invalidates_login_and_scope_and_rejects_form_edit(client, test_db, actors):
    done = await confirm(client, await upload(client, [teacher("depart-me")]))
    identifier = done["report"]["imported"][0]["id"]
    response = await client.put(P + f"/school/teachers/{identifier}", json={"display_name": "不可原地改名", "status": "active"})
    assert response.status_code == 409
    await lifecycle(client, "teachers", identifier, action="depart")
    assert (await test_db.get(User, UUID(identifier))).status == "departed"
    assert (await test_db.get(TeacherProfile, UUID(identifier))) is not None
    assert (await client.post("/api/v1/auth/login", json={"username": "depart-me", "password": "depart-me"})).status_code == 403


@pytest.mark.asyncio
async def test_original_embedded_templates_preflight_all_sample_rows(client, actors):
    await classroom(client)
    for kind, count in [("teachers", 4), ("students", 1)]:
        response = await client.get(P + f"/roster/{kind}/template")
        batch = await post(client, f"/roster/{kind}/imports", {"filename": "原始模板.xlsx", "content_base64": base64.b64encode(response.content).decode()})
        assert batch["report"]["ok"] and batch["report"]["total_rows"] == count


@pytest.mark.asyncio
async def test_mid_batch_failure_rolls_back_accounts_profiles_and_grants(client, test_db, actors, monkeypatch):
    from app.services import roster_workbench
    from app.core.errors import ApiError
    batch = await upload(client, [teacher("atomic-first"), teacher("atomic-second", "另一教师")])
    original = roster_workbench.grant_roles
    calls = 0

    async def fail_second(db, user, codes):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise ApiError(422, "SYNTHETIC_WRITE_FAILURE", "合成写入失败")
        return await original(db, user, codes)

    monkeypatch.setattr(roster_workbench, "grant_roles", fail_second)
    await confirm(client, batch, status=422)
    assert await test_db.scalar(select(func.count()).select_from(User).where(User.username.in_(["atomic-first", "atomic-second"]))) == 0
    assert await test_db.scalar(select(func.count()).select_from(TeacherProfile)) == 0
    assert (await client.get(P + f"/roster/imports/{batch['id']}")).json()["status"] == "ready"


@pytest.mark.asyncio
async def test_subject_and_grade_leader_grants_use_distinct_column_formats(client, test_db, actors):
    first = await classroom(client)
    await post(client, "/school/classes", {"external_class_id": "11.1", "name": "高中二年级1班"})
    imported = await confirm(client, await upload(client, [teacher("subject-leader", 教研组长负责年级科目="10.数学"), teacher("grade-leader", 年级长负责年级="10")]))
    for row, role in zip(imported["report"]["imported"], ["SUBJECT_LEADER", "GRADE_LEADER"]):
        profile = await test_db.get(TeacherProfile, UUID(row["id"]))
        assert role in profile.duties
        grants = (await test_db.scalars(select(TeachingAssignment).where(TeachingAssignment.teacher_user_id == profile.id))).all()
        assert len(grants) == 1 and str(grants[0].class_id) == first["id"]


@pytest.mark.asyncio
async def test_concurrent_confirmation_creates_one_batch(client, test_db, actors):
    if test_db.bind.dialect.name != "postgresql":
        pytest.skip("Concurrent school locks require PostgreSQL")
    import asyncio
    from sqlalchemy.ext.asyncio import async_sessionmaker
    from app.api.deps import AuthenticatedUser
    from app.db.models import School
    from app.schemas.roster import RosterConfirm
    from app.services.roster_workbench import confirm as confirm_service
    batch = await upload(client, [teacher("concurrent-account")])
    actor = AuthenticatedUser(actors["admin"].id, actors["school"].id, "qa-admin", ["SCHOOL_ADMIN"])
    factory = async_sessionmaker(test_db.bind, expire_on_commit=False)
    await test_db.commit()

    async def execute():
        async with factory() as db:
            await db.scalar(select(School.id).where(School.id == actor.school_id).with_for_update())
            return await confirm_service(db, actor, UUID(batch["id"]), RosterConfirm(expected_revision=batch["revision"]))

    first, second = await asyncio.gather(execute(), execute())
    assert first["report"] == second["report"]
    assert await test_db.scalar(select(func.count()).select_from(User).where(User.username == "concurrent-account")) == 1


@pytest.mark.asyncio
async def test_roster_permissions_tenancy_and_error_download(client, actors):
    batch = await upload(client, [teacher("valid")])
    for role in ("student", "teacher"):
        actors["switch"](role)
        assert (await client.get(P + "/roster/teachers/records")).status_code == 403
        assert (await client.get(P + f"/roster/imports/{batch['id']}/errors")).status_code == 403
    actors["switch"]("qa")
    await confirm(client, batch, status=403)
    actors["switch"]("foreign")
    assert (await client.get(P + f"/roster/imports/{batch['id']}")).status_code == 404
    await confirm(client, batch, status=404)
