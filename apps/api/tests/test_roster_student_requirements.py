"""Acceptance for the student additions required by the supplied 0922 HTML."""
import io
import base64
from copy import deepcopy
from uuid import UUID

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.api.deps import get_current_user
from app.core.security import get_password_hash
from app.db.models import AuditLog, QuestionScore, Student, StudentExamScore, StudentSubjectScore, User
from app.db.models.roster import RosterImport
from app.main import app
from app.services.roster_excel import STUDENT_COLUMNS, STUDENT_ERROR_COLUMNS
from test_education_workbench import actors, add_scores, post, setup_exam, P  # noqa: F401
from test_roster_workflows import confirm, lifecycle, student, teacher, upload, workbook_bytes


@pytest.mark.asyncio
@pytest.mark.parametrize("kind, row_factory, columns", [("students", student, STUDENT_COLUMNS), ("teachers", teacher, None)])
async def test_compact_roster_upload_returns_batch_reference_only(client, actors, kind, row_factory, columns):
    from app.services.roster_excel import TEACHER_COLUMNS

    columns = columns or TEACHER_COLUMNS
    rows = [row_factory(f"compact-{kind}-{index}") for index in range(120)]
    content = workbook_bytes(columns, [[row.get(field, "") for field in columns] for row in rows])
    response = await client.post(
        P + f"/roster/{kind}/imports?compact=true",
        json={"filename": "合成资料.xlsx", "content_base64": base64.b64encode(content).decode()},
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert set(result) == {"id", "kind", "status", "revision"}
    assert result["kind"] == kind and result["status"] == "ready"
    detail = await client.get(P + f"/roster/imports/{result['id']}")
    assert detail.status_code == 200
    assert detail.json()["report"]["total_rows"] == 120
    assert len(detail.content) > len(response.content) * 100
    assert detail.json()["rows"][0]["values"]["账号" if kind == "students" else "教师账号"] == f"compact-{kind}-0"


@pytest.mark.asyncio
async def test_compact_student_confirmation_keeps_full_detail_readback(client, actors):
    columns = STUDENT_COLUMNS
    row = student("compact-confirm-student")
    content = workbook_bytes(columns, [[row.get(field, "") for field in columns]])
    uploaded = await client.post(
        P + "/roster/students/imports?compact=true",
        json={"filename": "合成资料.xlsx", "content_base64": base64.b64encode(content).decode()},
    )
    assert uploaded.status_code == 200
    batch = uploaded.json()
    confirmed = await client.post(
        P + f"/roster/imports/{batch['id']}/confirm?compact=true",
        json={"expected_revision": batch["revision"]},
    )
    assert confirmed.status_code == 200, confirmed.text
    assert set(confirmed.json()) == {"id", "kind", "status", "revision"}
    assert confirmed.json()["status"] == "succeeded"
    detail = (await client.get(P + f"/roster/imports/{batch['id']}")).json()
    assert detail["report"]["success_count"] == 1
    assert len(detail["rows"]) == 1



@pytest.mark.asyncio
@pytest.mark.parametrize("state", ["", "休学", "退学", "毕业", "未知", " 正常 "])
async def test_student_import_rejects_non_normal_status_without_partial_writes(client, test_db, actors, state):
    batch = await upload(client, [student("valid-status-row"), student("invalid-status-row", 状态=state)], "students")
    assert batch["status"] == "invalid"
    assert batch["report"]["total_rows"] == 2
    assert len(batch["report"]["issues"]) == 1
    issue = batch["report"]["issues"][0]
    assert issue["row_number"] == 3
    assert issue["error_type"] == "状态非法" and issue["fields"] == "状态"
    assert "仅允许「正常」" in issue["message"]
    assert all(action in issue["suggestion"] for action in ("休学", "退学", "毕业", "学生列表"))
    await confirm(client, batch, status=409)
    assert await test_db.scalar(select(func.count()).select_from(User).where(User.username.in_(["valid-status-row", "invalid-status-row"]))) == 0


@pytest.mark.asyncio
async def test_student_error_workbook_merges_all_errors_for_a_row_into_eight_columns(client, actors):
    batch = await upload(client, [student("", "", 状态="休学"), student("valid-error-batch")], "students")
    assert batch["status"] == "invalid"
    assert len(batch["report"]["issues"]) == 1
    issue = batch["report"]["issues"][0]
    assert issue["fields"] == "姓名、账号、状态"
    assert issue["error_type"] == "必填缺失、状态非法"
    assert issue["name"] == issue["account"] == "（空）"
    response = await client.get(P + f"/roster/imports/{batch['id']}/errors")
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    rows = list(book.active.values)
    assert rows[0] == tuple(STUDENT_ERROR_COLUMNS)
    assert len(rows) == 2 and len(rows[1]) == 8
    assert rows[1][0] == "行2" and rows[1][2:5] == ("姓名、账号、状态", "（空）", "（空）")
    assert "休学" in rows[1][6]
    book.close()


@pytest.mark.asyncio
@pytest.mark.parametrize("mutation", ["strip_template_spaces", "add_header_space", "swap_columns"])
async def test_student_header_must_match_exact_names_and_order(client, test_db, actors, mutation):
    columns = STUDENT_COLUMNS.copy()
    if mutation == "strip_template_spaces":
        columns = [column.strip() for column in columns]
    elif mutation == "add_header_space":
        columns[4] += " "
    else:
        columns[4], columns[5] = columns[5], columns[4]
    before = await test_db.scalar(select(func.count()).select_from(User))
    # The invalid row would yield row errors if structure checking did not stop first.
    batch = await upload(client, [student("", "", 状态="休学")], "students", columns=columns)
    assert batch["status"] == "invalid" and batch["rows"] == []
    assert batch["report"]["total_rows"] == 0
    assert len(batch["report"]["issues"]) == 1
    issue = batch["report"]["issues"][0]
    assert issue["row_number"] is None and issue["error_type"] == "列结构不符"
    await confirm(client, batch, status=409)
    assert await test_db.scalar(select(func.count()).select_from(User)) == before


@pytest.mark.asyncio
async def test_student_confirm_rechecks_legacy_ready_batch_status_before_creating_accounts(client, test_db, actors):
    batch = await upload(client, [student("legacy-ready-valid"), student("legacy-ready-invalid")], "students")
    assert batch["status"] == "ready"
    record = await test_db.get(RosterImport, UUID(batch["id"]))
    rows = deepcopy(record.rows)
    rows[1]["values"]["状态"] = "退学"
    record.rows = rows
    # Reproduce a ready batch persisted by the old validator, before the new rule.
    await test_db.commit()
    result = await confirm(client, batch, status=409)
    assert result["error"]["code"] == "ROSTER_CHANGED"
    assert await test_db.scalar(select(func.count()).select_from(User).where(User.username.in_(["legacy-ready-valid", "legacy-ready-invalid"]))) == 0
    refreshed = (await client.get(P + f"/roster/imports/{batch['id']}")).json()
    assert refreshed["status"] == "invalid" and refreshed["revision"] == batch["revision"] + 1
    assert len(refreshed["report"]["issues"]) == 1
    assert refreshed["report"]["issues"][0]["fields"] == "状态"
    response = await client.get(P + f"/roster/imports/{batch['id']}/errors")
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    errors = list(book.active.values)
    assert errors[0] == tuple(STUDENT_ERROR_COLUMNS)
    assert len(errors) == 2 and errors[1][:3] == ("行3", "状态非法", "状态")
    assert "退学" in errors[1][6]
    book.close()


@pytest.mark.asyncio
async def test_student_import_keeps_optional_fields_without_extra_format_or_dictionary_checks(client, test_db, actors):
    source = student("学校自定义账号", "学生甲", **{"年级（1-12）": "", "班级号": "", "学号": "", "身份证": "不校验格式", "手机号": "131****7000", "email": "无需邮箱格式"})
    done = await confirm(client, await upload(client, [source], "students"))
    item = await test_db.get(Student, UUID(done["report"]["imported"][0]["id"]))
    assert item.class_id is None and item.status == "active"
    assert all(item.profile_fields[key] == value for key, value in source.items())


@pytest.mark.asyncio
async def test_student_without_class_dictionary_is_displayed_and_filterable(client, test_db, actors):
    batch = await upload(client, [student("unmapped-class-student", **{"年级（1-12）": "高中二年级", "班级号": "37班"})], "students")
    assert batch["report"]["ok"]
    done = await confirm(client, batch)
    identifier = done["report"]["imported"][0]["id"]
    assert (await test_db.get(Student, UUID(identifier))).class_id is None
    records = (await client.get(P + "/roster/students/records", params={"search": "unmapped-class-student"})).json()
    assert records["total"] == 1
    item = records["items"][0]
    assert "高中二年级" in item["class_name"] and "37班" in item["class_name"]
    for query in ("高中二年级", "37班"):
        matches = (await client.get(P + "/roster/students/records", params={"class_name": query})).json()
        assert [row["id"] for row in matches["items"]] == [identifier]
    assert (await client.get(P + "/roster/students/records", params={"class_name": "38班"})).json()["total"] == 0


@pytest.mark.asyncio
@pytest.mark.parametrize(("action", "target"), [("suspend", "suspended"), ("withdraw", "withdrawn")])
async def test_student_status_actions_keep_history_and_revoke_all_login_paths(client, test_db, actors, action, target):
    exam = await setup_exam(client, actors)
    scores = await add_scores(test_db, actors, exam)
    item = await test_db.get(Student, actors["student"].id)
    user = await test_db.get(User, item.user_id)
    user.password_hash = get_password_hash("student-history-secret")
    item.profile_fields = {"姓名": item.name, "状态": "正常", "年级（1-12）": "高中一年级", "班级号": "1班"}
    original_fields = deepcopy(item.profile_fields)
    original_password_version = user.password_version
    account, identifier = user.username, str(item.id)
    await test_db.commit()
    login = await client.post("/api/v1/auth/login", json={"username": account, "password": "student-history-secret"})
    assert login.status_code == 200, login.text
    tokens = login.json()

    assert (await lifecycle(client, "students", identifier, action=action))["status"] == target
    assert (await lifecycle(client, "students", identifier, action=action))["status"] == target
    assert item.status == user.status == target
    assert user.password_version == original_password_version + 1
    assert item.profile_fields == original_fields
    for model, score in zip((StudentExamScore, StudentSubjectScore, QuestionScore), scores):
        assert (await test_db.get(model, score.id)).student_id == item.id
    assert await test_db.scalar(select(func.count()).select_from(AuditLog).where(AuditLog.action == f"education.roster.{action}")) == 1
    records = (await client.get(P + "/roster/students/records", params={"search": account})).json()
    assert records["total"] == 1 and records["items"][0]["status"] == target
    app.dependency_overrides.pop(get_current_user)
    assert (await client.post("/api/v1/auth/login", json={"username": account, "password": "student-history-secret"})).status_code == 403
    assert (await client.get("/api/v1/auth/me", headers={"Authorization": "Bearer " + tokens["access_token"]})).status_code == 403
    assert (await client.post("/api/v1/auth/refresh", json={"refresh_token": tokens["refresh_token"]})).status_code == 401


@pytest.mark.asyncio
@pytest.mark.parametrize("action", ["suspend", "withdraw"])
async def test_student_status_actions_are_school_admin_only_and_never_apply_to_teachers(client, test_db, actors, action):
    identifier = str(actors["student"].id)
    for role in ("student", "teacher", "qa"):
        actors["switch"](role)
        await lifecycle(client, "students", identifier, action=action, status=403)
    actors["switch"]("foreign")
    await lifecycle(client, "students", identifier, action=action, status=404)
    actors["switch"]("admin")
    denied = await client.post(P + f"/roster/students/records/{identifier}/actions", json={"action": action, "confirmed": False})
    assert denied.status_code == 422
    assert (await test_db.get(Student, UUID(identifier))).status == "active"
    done = await confirm(client, await upload(client, [teacher("teacher-student-action")]))
    teacher_id = done["report"]["imported"][0]["id"]
    result = await lifecycle(client, "teachers", teacher_id, action=action, status=422)
    assert result["error"]["code"] == "ROSTER_ACTION_INVALID"
    assert (await test_db.get(User, UUID(teacher_id))).status == "active"
